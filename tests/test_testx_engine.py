"""TestX engine, rules, plan and export tests. No Streamlit needed."""
from itertools import combinations, product
from pathlib import Path

import openpyxl
import pytest

from tools.testx.engine import Engine, TooManyCombinations, bind_rules
from tools.testx.export import build_workbook
from tools.testx.model import BLANK, Cond, Factor, Rule
from tools.testx.plan import Settings, generate
from tools.testx.rules import Catalog, detect_rules, manual_rule
from tools.testx.source import read_source

FIX = Path(__file__).parent / "fixtures" / "testx_sample.xlsx"


def F(name, *vals, mandatory=True):
    return Factor(name, name, "n", tuple(vals), mandatory)


def R(rid, when, then):
    return Rule(rid, "manual", rid, (when,), (then,))


def engine(factors, rules):
    bound, skipped = bind_rules(factors, rules)
    assert not skipped
    return Engine(factors, bound)


def valid_full(factors, rules):
    return engine(factors, rules).full(10_000)


def pairs_of(case, keys):
    return {(a, case[a], b, case[b]) for a, b in combinations(keys, 2)}


# ------------------------------------------------------------------ engine
def test_pairwise_without_rules_covers_every_pair_and_is_small():
    fs = [F("A", "a1", "a2", "a3"), F("B", "b1", "b2", "b3"), F("C", "c1", "c2", "c3"), F("D", "d1", "d2", "d3")]
    cases, total, unreachable = engine(fs, []).pairwise()
    keys = [f.key for f in fs]
    covered = set().union(*(pairs_of(c, keys) for c in cases))
    assert unreachable == 0 and len(covered) == total == 54
    assert len(cases) < 81                      # fewer than all combinations (3^4)


def test_pairwise_matches_true_reachable_pairs_under_rules():
    fs = [F("Pkg", "S", "P", "X"), F("Cpe", "Forti", "Sonic"), F("Model", "F", "T80", "T270"), F("Opt", "y", "n")]
    rules = [
        R("R1", Cond("Pkg", frozenset({"S"})), Cond("Model", frozenset({"F"}))),
        R("R2", Cond("Cpe", frozenset({"Forti"})), Cond("Model", frozenset({"F"}))),
        R("R3", Cond("Cpe", frozenset({"Sonic"})), Cond("Model", frozenset({"F"}), negate=True)),
    ]
    eng = engine(fs, rules)
    cases, total, unreachable = eng.pairwise()
    keys = [f.key for f in fs]
    reachable = set().union(*(pairs_of(c, keys) for c in valid_full(fs, rules)))
    covered = set().union(*(pairs_of(c, keys) for c in cases))
    assert covered == reachable
    assert total - unreachable == len(reachable)
    assert all(eng.check.consistent(c) for c in cases)          # no case breaks a rule
    assert unreachable > 0


def test_full_respects_rules_and_limit():
    fs = [F("A", "1", "2"), F("B", "x", "y")]
    rules = [R("R1", Cond("A", frozenset({"1"})), Cond("B", frozenset({"x"})))]
    assert len(valid_full(fs, rules)) == 3                       # (1,y) is forbidden
    with pytest.raises(TooManyCombinations):
        engine(fs, []).full(3)


def test_blank_values_and_hidden_rule():
    fs = [F("Provider", "own", "other"), Factor("Id", "Id", "n", ("v1", BLANK), False)]
    rules = [R("R1", Cond("Provider", frozenset({"other"})), Cond("Id", frozenset({BLANK})))]
    cases = valid_full(fs, rules)
    assert {"Provider": "other", "Id": "v1"} not in cases
    assert {"Provider": "other", "Id": BLANK} in cases


def test_deterministic():
    fs = [F("A", "1", "2", "3"), F("B", "x", "y", "z"), F("C", "p", "q")]
    assert engine(fs, []).pairwise()[0] == engine(fs, []).pairwise()[0]


def test_negatives_break_their_rule():
    fs = [F("Pkg", "S", "P"), F("Model", "F", "T")]
    rules = [R("R1", Cond("Pkg", frozenset({"S"})), Cond("Model", frozenset({"F"})))]
    eng = engine(fs, rules)
    negs, why = eng.negatives()
    assert not why and len(negs) == 1
    b, case, key, bad, broken = negs[0]
    assert case == {"Pkg": "S", "Model": "T"} and key == "Model" and bad == "T" and broken == ["R1"]


def test_rule_that_cannot_bind_is_reported_not_crashing():
    fs = [F("A", "1", "2")]
    bound, skipped = bind_rules(fs, [R("R9", Cond("Nope", frozenset({"x"})), Cond("A", frozenset({"1"})))])
    assert not bound and skipped[0][0] == "R9"


# ------------------------------------------------------------------ rules from the workbook
@pytest.fixture(scope="module")
def sample():
    src = read_source(FIX.read_bytes())
    cat = Catalog.from_nodes(src.nodes)
    rules, text_only = detect_rules(src.blocks, cat)
    return src, rules, text_only


def test_reads_hierarchy_and_scenarios(sample):
    src, _, _ = sample
    assert src.product_name == "CYS_EdgeDefend"
    assert [n.name for n in src.nodes] == ["CYS_EdgeDefend", "CYS_EdgeDefend Security License"]
    assert src.order_types == ["CREATE", "MODIFY", "DELETE"]


def test_lookup_table_becomes_rules(sample):
    _, rules, _ = sample
    std = next(r for r in rules if r.source == "table" and "Package Type is Standard THEN CPE Model" in r.text)
    assert "Max Users is 50" in std.text and "Concurrent Users is 25" in std.text and std.enabled
    pro = next(r for r in rules if "Package Type is Pro THEN" in r.text and "CPE Model" in r.text)
    assert any("N/A" in n for n in pro.notes)                    # N/A cell ignored, with a note


def test_tick_matrix_becomes_rules_with_rest_no(sample):
    _, rules, _ = sample
    std = next(r for r in rules if "Package Type is Standard THEN Traffic Filtering" in r.text)
    assert "Traffic Filtering is Yes" in std.text and "Standard Reporting is Yes" in std.text
    pro = next(r for r in rules if "Package Type is Pro THEN Traffic Filtering" in r.text)
    assert "Traffic Filtering is No" in pro.text                 # blank cell + "rest ... No" text


def test_sentence_rules(sample):
    _, rules, _ = sample
    texts = [r.text for r in rules]
    assert "IF CPE Type is Fortinet THEN CPE Model is FortiF" in texts
    assert "IF CPE Type is SonicWall THEN CPE Model is not FortiF" in texts
    assert any("CPE Type is not Fortinet THEN License Type is (blank)" in t for t in texts)


def test_rule_with_unusable_attribute_is_switched_off_and_explained(sample):
    _, rules, _ = sample
    hidden = next(r for r in rules if "[hidden]" in r.text)
    assert not hidden.enabled and hidden.status == "Needs review"
    assert "no list of values" in hidden.issues[0]


def test_component_rules_are_text_only(sample):
    _, _, text_only = sample
    assert any("Component-level" in t.reason for t in text_only)


def test_manual_rule():
    r = manual_rule(1, Cond("A", frozenset({"1"})), Cond("B", frozenset({"x"}), negate=True))
    assert r.id == "M01" and "THEN B is not x" in r.text and r.enabled


# ------------------------------------------------------------------ plan + excel
def test_plan_per_component_and_together(sample):
    src, rules, _ = sample
    per = generate(src, rules, Settings(scope="component"), src.product_name)
    assert [g.name for g in per.groups] == ["CYS_EdgeDefend", "CYS_EdgeDefend Security License"]
    g = per.groups[0]
    assert g.full_size > 100_000 and len([c for c in g.cases if c.kind == "Positive"]) < 100
    keys = [f.key for f in g.factors]
    eng = Engine(g.factors, bind_rules(g.factors, rules)[0])
    assert all(eng.check.consistent(c.values) for c in g.cases if c.kind == "Positive")
    together = generate(src, rules, Settings(scope="together"), src.product_name)
    assert [g.name for g in together.groups] == ["All components"]
    assert any(f.name == "License Type" for f in together.groups[0].factors)
    # the 'License Type only when Fortinet' rule can only be tested when both attributes are together
    lic = next(r.id for r in rules if "License Type is (blank)" in r.text)
    assert lic in together.groups[0].rules_applied and lic not in per.groups[0].rules_applied


def test_full_technique_too_big_is_refused(sample):
    src, rules, _ = sample
    with pytest.raises(TooManyCombinations):
        generate(src, rules, Settings(technique="full", full_limit=5), src.product_name)


def test_disabled_rule_is_not_enforced(sample):
    src, rules, _ = sample
    off = [Rule(r.id, r.source, r.text, r.when, r.then, enabled=False) for r in rules]
    plan = generate(src, off, Settings(negatives=True), src.product_name)
    assert plan.negative == 0 and all(not g.rules_applied for g in plan.groups)


def test_excel_export_is_a_usable_uat_sheet(sample, tmp_path):
    src, rules, _ = sample
    plan = generate(src, rules, Settings(vary_order_type=True), src.product_name)
    out = tmp_path / "plan.xlsx"
    out.write_bytes(build_workbook(plan, "testx_sample.xlsx", "BPM team"))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Rules"
    ws = wb[wb.sheetnames[1]]
    headers = [c.value for c in ws[5]]
    assert headers[:2] == ["Test ID", "Type"] and "Order Type" in headers
    assert headers[-7:] == ["Rules", "Expected Result", "Actual Result", "Status", "Tested By", "Test Date", "Remarks"]
    status_col = headers.index("Status") + 1
    assert all(ws.cell(r, status_col).value == "Not Run" for r in range(6, ws.max_row + 1))
    assert ws.freeze_panes == "C6" and ws.auto_filter.ref
    assert any(dv.formula1 and "Pass" in dv.formula1 for dv in ws.data_validations.dataValidation)
    ids = [ws.cell(r, 1).value for r in range(6, ws.max_row + 1)]
    assert len(ids) == len(set(ids))
    summary = wb["Summary"]
    assert any(c.value == "Total" for row in summary.iter_rows() for c in row)
