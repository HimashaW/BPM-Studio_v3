from docx import Document
from tempfile import NamedTemporaryFile
import zipfile

from models.schemas import Task
from render.task_flow import prepare_flow_tasks, add_task_flowchart


def make_task(seq, name, system, work_group):
    return Task(
        seq=str(seq),
        task_name=name,
        it_system=system,
        work_group=work_group,
        source_sheet="6.3 OSS SO Task List ",
        source_row=1,
    )


def test_flow_rules_filter_and_lane_fallback():
    tasks = [
        make_task(1, "Initiate Service Order", "CRM", "CRM"),
        make_task(2, "System Validation", "CRM", "System"),
        make_task(3, "Collect CPE", "RTOM-DSC-PROV", ""),
        make_task(4, "External Confirmation", "CRM", "External"),
        make_task(5, "Approve SO", "OSS", "Operations"),
    ]

    prepared = prepare_flow_tasks(tasks)

    assert [t.seq for t in prepared] == ["1", "3", "5"]
    assert [t.lane for t in prepared] == ["CRM", "RTOM-DSC-PROV", "Operations"]


def test_flowchart_is_inserted_as_editable_vml():
    doc = Document()
    tasks = [
        make_task(1, "Initiate Service Order", "CRM", "CRM"),
        make_task(2, "Approve SO", "OSS", "Operations"),
        make_task(3, "Collect CPE", "RTOM-DSC-PROV", ""),
    ]

    summary = add_task_flowchart(doc, tasks)
    assert summary["editable_independent_shapes"] is True
    assert summary["eligible_tasks"] == 3
    assert summary["panels"] == 1
    assert summary["lanes"] == ["CRM", "Operations", "RTOM-DSC-PROV"]

    xml = doc._body._element.xml
    assert "v:group" not in xml
    assert "D9ECFA" in xml
    assert "5B9BD5" in xml
    assert "v:roundrect" not in xml
    assert "v:oval" in xml
    assert "v:rect" in xml
    assert "Initiate" in xml and "Service" in xml and "Order" in xml
    assert "1. Initiate" not in xml
    assert "Approve" in xml and "SO" in xml
    assert "Collect" in xml and "CPE" in xml
    assert "External" not in xml


def test_duplicate_sequences_are_kept_and_connected_from_first_task():
    tasks = [
        make_task(10, "RESERVE CPE", "OSS", "Provisioning"),
        make_task("10-ONBOARD CPE", "ONBOARD CPE", "OSS", "Provisioning"),
        make_task(20, "ACTIVATE SERVICE", "OSS", "Activation"),
    ]
    prepared = prepare_flow_tasks(tasks)
    assert [t.seq for t in prepared] == ["10", "10-ONBOARD CPE", "20"]

    doc = Document()
    add_task_flowchart(doc, tasks)
    with NamedTemporaryFile(suffix=".docx") as f:
        doc.save(f.name)
        with zipfile.ZipFile(f.name) as z:
            xml = z.read("word/document.xml").decode("utf-8")
    assert "RESERVE" in xml and "CPE" in xml
    assert "ONBOARD" in xml and "CPE" in xml
    assert "ACTIVATE" in xml and "SERVICE" in xml
    # Duplicate-sequence connector is explicitly emitted as an orthogonal
    # vertical route, and sequence continuation is emitted separately.
    assert "duplicate_connector_1_0_1" in xml
    assert "sequence_connector_1_0" in xml


def test_excluded_task_prefixes_are_omitted():
    tasks = [
        make_task(1, "OM SR_REQ 123", "OSS", "Provisioning"),
        make_task(2, "WAIT FOR AM RESPOND", "OSS", "Provisioning"),
        make_task(3, "ADD REWORK ACTIVITY1", "OSS", "Provisioning"),
        make_task(4, "SET AM INVOLV BLANK", "OSS", "Provisioning"),
        make_task(5, "RECONFIRM W/CUSTOMER", "OSS", "Provisioning"),
        make_task(6, "KEEP THIS TASK", "OSS", "Provisioning"),
    ]
    prepared = prepare_flow_tasks(tasks)
    assert [t.task_name for t in prepared] == ["KEEP THIS TASK"]

    doc = Document()
    add_task_flowchart(doc, tasks)
    with NamedTemporaryFile(suffix=".docx") as f:
        doc.save(f.name)
        with zipfile.ZipFile(f.name) as z:
            xml = z.read("word/document.xml").decode("utf-8")
    assert "KEEP" in xml and "THIS" in xml and "TASK" in xml
    for name in ["OM SR_REQ", "WAIT FOR AM RESPOND", "ADD REWORK ACTIVITY1",
                 "SET AM INVOLV BLANK", "RECONFIRM W/CUSTOMER"]:
        assert name not in xml
