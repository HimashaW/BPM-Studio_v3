"""Scenario-style TestX tests. Uses tests/fixtures/sase_sample.xlsx (the SASE workbook, reduced to the sheets TestX reads)."""
from pathlib import Path

import openpyxl
import pytest

from tools.testx.export_cases import build_cases_workbook
from tools.testx.scenarios import build_scenarios, coverage_notes, short_order_type
from tools.testx.source import read_source

FIX = Path(__file__).parent / "fixtures" / "sase_sample.xlsx"
pytestmark = pytest.mark.skipif(not FIX.exists(), reason="sase_sample.xlsx fixture missing")


@pytest.fixture(scope="module")
def src():
    return read_source(FIX.read_bytes())


@pytest.fixture(scope="module")
def scen(src):
    return build_scenarios(src)


def test_short_order_type_names():
    assert short_order_type("Comprehensive Task List for D_SLT SASE  - Order Type  CREATE SO") == "CREATE"
    assert short_order_type("Comprehensive Task List for X - Order Type  MODIFY-SERVICE SO") == "MODIFY-SERVICE"
    assert short_order_type("Comprehensive Task List for X - Order Type DELETE SERVICE SO") == "DELETE"
    assert short_order_type("Comprehensive Task List for X - CREATE Order Type") == "CREATE"


def test_facts_are_read(src):
    assert src.product_name == "D_SLT SASE" and len(src.sections) == 5
    assert src.extra.task_sheet.strip().startswith("6.2")
    assert len(src.extra.pricing) > 20 and src.extra.faults and src.extra.requirements and src.extra.modify_rows
    assert not [w for w in src.warnings if "5.1" in w]


def test_expected_scenarios_exist(scen):
    keys = {s.key for s in scen}
    for k in ("attr:catalogue", "map:mandatory", "cond:user count", "child:user class", "servcat:service category",
              "modify:MODIFY-SERVICE", "fault:d slt sase", "flow:CREATE", "flow:MODIFY-SERVICE", "flow:DELETE", "flow:SUSPEND",
              "flow:RESUME", "task:cancel:CREATE", "task:confirm:CREATE", "task:test:CREATE"):
        assert k in keys, k
    sr = [s for s in scen if s.key.startswith("sr:")]
    assert len(sr) == 5                                      # REV1/2/3 on CREATE, REV2/3 on MODIFY-SERVICE
    assert {s.order_type for s in sr} == {"CREATE", "MODIFY-SERVICE"}


def test_recommended_set_is_lean_and_unique(scen):
    core = [s for s in scen if s.core]
    assert 18 <= len(core) <= 30 < len(scen)
    assert len({s.key for s in scen}) == len(scen)
    assert all(s.text.strip() and s.basis for s in scen)
    assert not any(s.core for s in scen if s.category == "Special requirements")      # opt-in, so the plan stays short


def test_sr_scenario_is_derived_from_the_task(scen):
    s = next(x for x in scen if x.key == "sr:CREATE:om sr req sol rev1")
    assert "'CONFIGURE TENANT'" in s.text and "CUSTOMER NOT READY" in s.text
    assert "WAIT FOR AM RESPOND1" in s.oss and "hidden task" not in s.oss
    assert s.crm == "SR is passed to CRM with the correct attributes."


def test_conditional_attribute_and_pricing(scen):
    c = next(x for x in scen if x.key == "cond:user count")
    assert "More than 1000 Users" in c.crm and "greater than 1000" in c.crm
    price = [x for x in scen if x.category == "Pricing"]
    assert len(price) >= 3 and sum(x.core for x in price) == 1
    assert "10000000" in next(x for x in price if x.core).crm


def test_notes_name_order_types_without_task_lists(src):
    assert any("TRANSFER" in n for n in coverage_notes(src))


def test_workbook_without_task_sheet_still_works():
    wb = openpyxl.load_workbook(FIX)
    del wb["6.2 OSS SO Task List "]
    import io
    buf = io.BytesIO()
    wb.save(buf)
    s2 = read_source(buf.getvalue())
    sc = build_scenarios(s2)
    assert sc and not any(x.key.startswith(("flow:", "sr:")) for x in sc)
    assert any("Task list" in w for w in s2.warnings)


def test_excel_layout_matches_the_manual_format(scen, tmp_path):
    chosen = [s for s in scen if s.core]
    out = tmp_path / "cases.xlsx"
    out.write_bytes(build_cases_workbook(chosen, "D_SLT SASE", "x.xlsx", "BPM team", ["a note"]))
    wb = openpyxl.load_workbook(out)
    ws = wb["Test Cases"]
    assert [ws.cell(1, c).value for c in range(1, 8)] == ["Test Case ID", "Service Type", "Order Type", "IT System", "Product",
                                                           "Test Scenario", "Expected result in CRM/OSS/BSS"]
    assert [ws.cell(2, c).value for c in (7, 8, 9)] == ["CRM", "OSS", "BSS"] and ws.cell(1, 10).value == "Status"
    assert "G1:I1" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.max_row == len(chosen) + 2 and ws.cell(3, 1).value == 1 and ws.cell(3, 10).value == "Not Run"
    assert ws.freeze_panes == "A3" and any("Pass" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)
    assert wb.sheetnames == ["Test Cases", "Summary"]
