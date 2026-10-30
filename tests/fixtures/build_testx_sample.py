"""Builds tests/fixtures/testx_sample.xlsx - a small workbook shaped like a real product configuration
document (sheet 2 hierarchy, sheet 3 special requirements with tables, sheet 5.1 order scenarios).
Run:  python tests/fixtures/build_testx_sample.py
"""
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
base = openpyxl.load_workbook(HERE / "special_requirements_two_level_header.xlsx")   # real-shaped sheet 3
wb = base

ws = wb.create_sheet("2. Product Hierarchy")
ws.append(["A", "B", "Entity Type", "Entity Name", "Mandatory?", "Comments"])


def node(kind, name):
    ws.append([None, None, kind, name, None, None])


def attr(name, mandatory, values=()):
    ws.append([None, None, "Attribute", name, mandatory, None])
    for v in values:
        ws.append([None, None, "Value", v, None, None])


node("Product", "CYS_EdgeDefend")
attr("Package Type", "Yes", ["Standard", "Pro", "Premium", "Advanced Pro", "Power", "Ultra", "Max"])
attr("CPE Type", "Yes", ["Fortinet", "SonicWall"])
attr("CPE Model", "Yes", ["FortiF", "TZ80", "TZ270", "TZ370"])
attr("Max Users", "Yes", ["50", "100", "200"])
attr("Concurrent Users", "Yes", ["25", "50", "75", "100"])
attr("Internet Provider", "Yes", ["Own Network", "Other Operator"])
attr("Internet Provider ID", "No")
for feature in ["Traffic Filtering", "Application Visibility & Control", "Network Address Translation (NAT)",
                "VPN Support", "Layer 2-4 Protection", "Standard Reporting", "Advanced Web Filtering"]:
    attr(feature, "Yes", ["Yes", "No"])
node("Product Component", "CYS_EdgeDefend Security License")
attr("License Type", "No", ["Basic", "Advanced"])

sc = wb.create_sheet("5.1 MDM_Order Scenarios")
sc.append(["Service Type", "Scenario ID", "Order Sub Type", "Condition", "Root Product Action", "Order Type", "Comments"])
sc.append(["Edge", 1, "New", "Standard", "Add", "CREATE", None])
sc.append(["Edge", 2, "Change", "Standard", "Modify", "MODIFY", None])
sc.append(["Edge", 3, "Cease", "Standard", "Delete", "DELETE", None])

wb.save(HERE / "testx_sample.xlsx")
print("written", HERE / "testx_sample.xlsx")
