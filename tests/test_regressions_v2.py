"""Regression tests: two-level table headers, R10/R11 not a table,
LOV attribute name written once, flowchart images."""
import os, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from extractors.special_requirements import extract_business_rule_blocks
from render import components as C
from render.task_flow_image import render_flow_images
from models.schemas import Task

FIX = os.path.join(os.path.dirname(__file__), "fixtures", "special_requirements_two_level_header.xlsx")


def test_r05_table_keeps_all_columns_and_group_header():
    blocks = extract_business_rule_blocks(FIX)
    r05 = [b for b in blocks if b.text.startswith("The attribute value of the following feature")][0]
    t = r05.tables[0]
    assert t.headers == ["Feature Attribute", "Standard", "Pro", "Premium", "Advanced Pro", "Power", "Ultra", "Max"]
    assert t.group_headers == [["", 1], ['Attribute values of "Package Type"', 7]]
    assert len(t.rows) == 16


def test_r10_r11_are_plain_requirements_not_tables():
    blocks = extract_business_rule_blocks(FIX)
    by = {b.text[:25]: b for b in blocks}
    r11 = [b for b in blocks if "Advanced Reporting" in b.text][0]
    assert r11.tables == []
    assert any(b.text.startswith("This value shall be the default") for b in blocks)
    assert all(not b.tables for b in blocks if "feature attributes" not in b.text and "LOVs" not in b.text)


def test_lov_attribute_name_written_once():
    doc = Document()
    C.add_attribute_values_table(doc, [("CPE Type", ["Fortinet", "SonicWall", "Other"])])
    cell = doc.tables[0].rows[1].cells[0]
    assert cell.text.strip() == "CPE Type"


def test_flow_images_created_per_order_type():
    tasks = [Task(seq="00-1", task_name="INITIATE", it_system="CRM", source_sheet="s", source_row=1),
             Task(seq="10", task_name="RESERVE CPE", work_group="PROV", it_system="OSS", source_sheet="s", source_row=2),
             Task(seq=None, flow_seq="10", task_name="10-ONBOARD", work_group="MS", it_system="OSS", source_sheet="s", source_row=3)]
    with tempfile.TemporaryDirectory() as d:
        paths, _warn = render_flow_images(tasks, d, "CREATE")
        assert paths and all(os.path.getsize(p) > 1000 for p in paths)


def _t(seq, name, wg=None, it="OSS", cond=None, flow_seq=None, row=1):
    return Task(seq=seq, flow_seq=flow_seq, task_name=name, work_group=wg, it_system=it,
                task_condition=cond, source_sheet="s", source_row=row)


def test_strict_sequence_rework_excluded_and_backwards_row_reported():
    from render.task_flow import prepare_flow_tasks_ex, condition_label
    tasks = [_t("00-1", "INITIATE SERVICE ORDER", it="CRM"), _t("00-3", "APROVE SO", it="CRM"),
             _t("10", "RESERVE CPE", "<RTOM>-DSC-PROV"), _t(None, "10-ONBOARD CPE", "CEN-CSC-MS", flow_seq="10"),
             _t("30", "INSTALL CPE", "<RTOM>-DSC-PROV"), _t("50", "CONFIGURE CPE", "ISEC-OPR CEN-CSC-MS"),
             _t("10", "INSTALL CPE", "<RTOM>-DSC-PROV", row=11),            # list goes backwards
             _t("53", "ADD REWORK ACTIVITY1", "CEN-CSC-MS"), _t("63", "ADD REWORK ACTIVITY3", "CEN-CSC-MS"),
             _t("80", "CLOSE SERVICE ORDER", "AUTO")]
    flow, warns = prepare_flow_tasks_ex(tasks)
    assert [f.seq for f in flow] == ["00-1", "00-3", "10", "10", "30", "50", "80"]
    assert not any("REWORK" in f.task_name for f in flow)
    assert len(warns) == 1 and "Seq 10" in warns[0]
    assert flow[3].lane == "ISEC-OPR"            # CEN-CSC-MS shares the ISEC-OPR lane
    assert condition_label("if INTERNET PROVIDER = 'SLT\"") == "Internet Provider SLT?"
    assert condition_label("if [CPE MODEL '<New Value>' <> '<Old Value>'") == "CPE Model changed?"


def test_delete_reference_lanes_and_flow():
    """DELETE (reference chart): CRM-system tasks share the CRM lane; CEN-CSC-MS keeps its name."""
    from render.task_flow_layout import build_layouts
    tasks = [_t("00-1", "INITIATE SERVICE ORDER", "<SECTION HANDLED BY>-FO", it="CRM"),
             _t("00-2", "APPROVE SO", "<SECTION HANDLED BY>-MGR", it="CRM"),
             _t("10", "DEACTIVATE SVS", "CEN-CSC-MS"), _t("20", "COLLECT CPE", "<RTOM>-DSC-PROV"),
             _t("30", "CLOSE SO", "AUTO")]
    (lay,), warns = build_layouts(tasks, "DELETE", "CYS_EdgeDefend")
    assert lay.lanes == ["CRM", "<RTOM>-DSC-PROV", "CEN-CSC-MS", "AUTO"]
    assert lay.title == "CYS_EdgeDefend-DELETE workflow"
    labels = [n.label.replace("\n", " ") for n in lay.nodes]
    assert labels[:5] == ["Initiate Service Order", "Approve SO", "Deactivate SVS", "Collect CPE", "Close SO"]
    assert len(lay.edges) == 6        # start->..->end, one arrow per step


def test_pptx_export_has_native_shapes_and_glued_connectors():
    import tempfile
    from pptx import Presentation
    from render.task_flow_pptx import build_flowcharts_pptx

    class Sec:
        order_type = "CREATE"
        tasks = [_t("00-1", "INITIATE SERVICE ORDER", it="CRM"), _t("10", "RESERVE CPE", "<RTOM>-DSC-PROV"),
                 _t(None, "10-ONBOARD CPE", "CEN-CSC-MS", flow_seq="10"),
                 _t("20", "CHECK LINK READINESS", "AUTO", cond="if INTERNET PROVIDER = 'SLT'"),
                 _t("30", "INSTALL CPE", "<RTOM>-DSC-PROV")]
    with tempfile.TemporaryDirectory() as d:
        info = build_flowcharts_pptx([Sec], d + "/f.pptx", "P")
        prs = Presentation(info["path"])
        xml = prs.slides[0].shapes._spTree.xml
        assert info["slides"] == 1 and "flowChartDecision" in xml
        assert "<a:stCxn" in xml and "<a:endCxn" in xml          # connectors glued to shapes
        assert xml.count('type="triangle"') >= 6                  # arrow heads present


def test_web_service_helpers():
    import app_service as svc
    wb = os.path.join(os.path.dirname(__file__), "fixtures", "special_requirements_two_level_header.xlsx")
    info = svc.inspect_workbook(open(wb, "rb").read())
    assert any(c["prefix"] == "3." and c["found"] for c in info["checks"])
    assert svc.default_product_name("Product_Configuration_Document_v_1_3_MS-Edge_Defender_02Apr2026.xlsx") == "MS-Edge Defender"
    assert svc.safe_name("My Product!") == "My_Product"


def _example_tasks():
    C1 = "if [CPE MODEL '<New Value>' <> '<Old Value>'"
    return [_t("00-1", "INITIATE SERVICE ORDER", it="CRM"), _t("00-3", "APROVE SO", it="CRM"),
            _t("10", "WA-MODIFY CPE", "<RTOM>-DSC-PROV", cond=C1),
            _t(None, "WA-20-CONNECT CPE ", "CEN-CSC-MS", flow_seq="10"),
            _t("20", "COLLECT CPE", "<RTOM>-DSC-PROV", cond=C1), _t("20", "RECONFIG CPE", "CEN-CSC-MS"),
            _t("30", "MODIFY LICENSE", "CEN-CSC-MS", cond="if [LICENSE TYPE '<New Value>' <> '<Old Value>'"),
            _t("30", "ADD/ REM FEATURE", "CEN-CSC-MS", cond="if [ADVANCED REPORTING REQ.'<New Value>' <> '<Old Value>'"),
            _t("40", "CONFIRM  W/ CUSTOMER", "CEN-CSC-MS"), _t("50", "CLOSE SERVICE ORDER", "AUTO")]


def test_work_activity_vs_parallel_tasks():
    from render.task_flow import prepare_flow_tasks
    from render.task_flow_layout import build_layouts, classify, group_tasks
    flow = prepare_flow_tasks(_example_tasks())
    groups = {g[0].seq: g for g in group_tasks(flow)}
    mains10, was10 = classify(groups["10"])                  # WA-MODIFY CPE + WA-20-CONNECT CPE
    assert [t.label for t in mains10] == ["WA-Modify CPE"] and [t.label for t in was10] == ["WA-Connect CPE"]
    assert len(classify(groups["20"])[0]) == 2 and not classify(groups["20"])[1]      # parallel
    assert len(classify(groups["30"])[0]) == 2
    lays, _ = build_layouts(_example_tasks(), "MODIFY-PACKAGE", "P")
    both = sum(1 for l in lays for e in l.edges if e.heads == "both")
    assert both == 1                                          # only the Work Activity pair is double-headed
    assert all(l.lanes == lays[0].lanes for l in lays) and len(lays[0].lanes) == 4    # every part shows ALL lanes


def test_diagram_fits_one_page_or_splits_with_all_lanes():
    from render.task_flow_layout import build_layouts, H_AVAIL
    tasks = [_t(f"{i*10:02d}", f"STEP {i}", ["CRM", "AUTO", "<RTOM>-DSC-PROV"][i % 3], it="OSS") for i in range(1, 25)]
    lays, _ = build_layouts(tasks, "BIG", "P")
    assert len(lays) >= 2 and all(l.height * l.scale <= H_AVAIL + 1e-6 for l in lays)
    assert all(l.lanes == lays[0].lanes for l in lays)
    small, _ = build_layouts(tasks[:5], "SMALL", "P")
    assert len(small) == 1
