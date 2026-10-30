"""Editable Word swimlane task-flow diagrams.

The diagram is built with legacy Word VML shapes inside the DOCX.  VML is
used intentionally here because python-docx does not expose a public API for
editable floating Word shapes/connectors.  Word opens these shapes as normal
editable drawing objects, so users can move/edit task boxes and connectors.

Rules implemented from the BPM task-table specification:
- Tasks are ordered by Task Seq.
- Work Group is the swimlane.
- If Work Group is blank, IT System becomes the swimlane.
- Work Group values External/System are excluded from the diagram.
- The source task table is not modified; filtering only affects the diagram.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from typing import Iterable

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from models.schemas import Task

VML_NS = "urn:schemas-microsoft-com:vml"

def _v(tag: str):
    """Create a VML element without relying on python-docx's nsmap."""
    return etree.Element(f"{{{VML_NS}}}{tag}")

# Word/VML coordinate system.  The group is displayed at this physical size.
GROUP_WIDTH_IN = 6.25
GROUP_HEIGHT_IN = 5.35
COORD_WIDTH = 625
COORD_HEIGHT = 535
# VML shapes are emitted as independent siblings (not a v:group).
# 100 VML layout units correspond to 1 inch in the original layout.
COORD_TO_PT = (GROUP_WIDTH_IN * 72.0) / COORD_WIDTH

# Light-blue palette for the editable diagram.
LANE_FILL = "EEF6FC"
LANE_HEADER_FILL = "CFE8F8"
TASK_FILL = "D9ECFA"
TASK_STROKE = "5B9BD5"
START_END_FILL = "EAF5FC"
START_END_STROKE = "5B9BD5"
CONNECTOR_COLOR = "5B9BD5"

# Keep the diagram readable on a portrait A4 page.  A flow with more tasks is
# split into panels; sequence continues naturally from one panel to the next.
MAX_TASKS_PER_PANEL = 8

LANE_HEADER_H = 42
LANE_PAD_X = 5
# All workflow symbols use one consistent size within a panel. Swimlane
# backgrounds/headers are intentionally larger because they define the lanes.
FLOW_SHAPE_W = 90
FLOW_SHAPE_H = 36
TOP_PAD = 58
BOTTOM_PAD = 28
STEP_Y = 52

import itertools
_PANEL_UID = itertools.count(1)

import os as _os

_DEFAULT_RULES = {
    "excluded_work_groups": ["external", "system"],
    "excluded_task_prefixes": ["OM SR_REQ", "WAIT FOR AM RESPOND", "ADD REWORK ACTIVITY",
                               "SET AM INVOLV BLANK", "RECONFIRM W/CUSTOMER"],
    "lane_aliases": {}, "lane_merges": [], "it_system_lanes": [],
    "lane_order": [], "exit_task_overrides": [],
    "work_activity_pattern": r"^WA\s*-",
    "acronyms": ["CPE", "SO", "ID", "SLT", "CRM", "OSS", "SVS", "SR", "AM", "OR", "WA"],
}
_RULES_CACHE: dict | None = None


def load_flow_rules(reload: bool = False) -> dict:
    """Read mapping/flow_rules.yaml (falls back to built-in defaults)."""
    global _RULES_CACHE
    if _RULES_CACHE is not None and not reload:
        return _RULES_CACHE
    rules = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in _DEFAULT_RULES.items()}
    path = _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))),
                         "mapping", "flow_rules.yaml")
    try:
        import yaml
        with open(path, encoding="utf-8") as f:
            loaded = yaml.safe_load(f) or {}
        for k in rules:
            if k in loaded and loaded[k] is not None:
                rules[k] = loaded[k]
    except Exception:
        pass
    _RULES_CACHE = rules
    return rules


EXCLUDED_WORK_GROUPS = set(x.lower() for x in load_flow_rules()["excluded_work_groups"])
# Optional semantic shape hints (kept for backward compatibility with the
# editable-VML renderer; the image renderer takes decisions from the Excel
# Task Condition column instead).
TASK_SHAPE_OVERRIDES: dict[tuple[str, str], str] = {}
# Task names beginning with any of these values are omitted from the diagram.
EXCLUDED_TASK_PREFIXES = tuple(load_flow_rules()["excluded_task_prefixes"])


@dataclass(frozen=True)
class FlowTask:
    seq: str
    task_name: str
    it_system: str
    work_group: str
    lane: str
    condition: str = ""
    label: str = ""
    is_wa: bool = False          # Work Activity: task name starts with "WA-"


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _natural_key(value: str):
    """Natural sort: 2 before 10, while preserving mixed/text sequences."""
    parts = re.split(r"(\d+(?:\.\d+)*)", _clean(value))
    key = []
    for part in parts:
        if not part:
            continue
        if re.fullmatch(r"\d+(?:\.\d+)*", part):
            key.append((0, tuple(int(x) for x in part.split("."))))
        else:
            key.append((1, part.lower()))
    return tuple(key)


def _norm_ws(value: str) -> str:
    return " ".join(_clean(value).split())


def title_label(text: str) -> str:
    """'CONFIRM  W/ CUSTOMER' -> 'Confirm W/ Customer' (acronyms stay upper)."""
    acr = {a.upper() for a in load_flow_rules()["acronyms"]}
    def fix(tok: str) -> str:
        def word(m):
            w = m.group(0)
            if w.upper() in acr or len(w) == 1:
                return w.upper()
            return w[:1].upper() + w[1:].lower()
        return re.sub(r"[A-Za-z]+", word, tok)
    return " ".join(fix(t) for t in _norm_ws(text).split(" "))


def is_work_activity(task_name: str) -> bool:
    """Work Activity = task name starts with 'WA-' (pattern in flow_rules.yaml)."""
    return re.match(load_flow_rules()["work_activity_pattern"], _norm_ws(task_name), flags=re.I) is not None


def display_name(task_name: str) -> str:
    """Box text. 'WA-20-CONNECT CPE' -> 'WA-Connect CPE'; '10-ONBOARD CPE' -> 'Onboard CPE'.
    ALL-CAPS names are shown in Title Case (acronyms stay upper)."""
    name = _norm_ws(task_name)
    m = re.match(r"^(WA)\s*-\s*(?:\d+\s*-\s*)?(.*)$", name, flags=re.I)
    if m:
        name = "WA-" + m.group(2)
    else:
        name = re.sub(r"^\s*\d+\s*-\s*", "", name)
    return title_label(name) if name.upper() == name else name


def lane_title(lane: str) -> str:
    lane = _norm_ws(lane)
    # A plain one-word lane such as AUTO -> Auto; names with symbols stay as written.
    return lane.capitalize() if re.fullmatch(r"[A-Za-z]+", lane) and lane.isupper() and lane.upper() not in {
        a.upper() for a in load_flow_rules()["acronyms"]} and len(lane) > 3 else lane


def condition_label(cond: str) -> str:
    """Short question for a decision diamond, built from the Excel Task Condition.

    "if INTERNET PROVIDER = 'SLT'"                  -> "Internet Provider SLT?"
    "if [CPE MODEL '<New Value>' <> '<Old Value>'"  -> "CPE Model changed?"
    "if [CPE MODEL '<New Value>' = '<Old Value>'"   -> "CPE Model unchanged?"
    Anything more complex falls back to the cleaned condition text.
    """
    c = _norm_ws(cond)
    c = re.sub(r"^if\s+", "", c, flags=re.I).strip().strip("[] ").strip()
    m = re.match(r"^(.*?)\s*'?<New Value>'?\s*(<>|!=|=)\s*'?<Old Value>'?\s*$", c, flags=re.I)
    if m:
        attr = title_label(m.group(1).strip(" '"))
        return f"{attr} {'unchanged' if m.group(2) == '=' else 'changed'}?"
    m = re.match(r"^'?([^'=<>]+?)'?\s*=\s*['\"]?([^'\"=<>]+?)['\"]?\s*$", c)
    if m and " and " not in c.lower() and " or " not in c.lower():
        return f"{title_label(m.group(1))} {m.group(2).strip()}?"
    short = c if len(c) <= 48 else c[:47].rstrip() + "…"
    return short + "?"


def _canon_lane(lane: str, aliases: dict) -> str:
    key = _norm_ws(lane).casefold()
    for src, dst in aliases.items():
        if _norm_ws(src).casefold() == key:
            return _norm_ws(dst)
    return _norm_ws(lane)


def prepare_flow_tasks_ex(tasks: Iterable[Task]):
    """Apply the flowchart rules. Returns (flow_tasks, warnings).

    * Rows are used in Excel order; a blank Seq (vertically merged cell)
      inherits the merged value, so it belongs to the same sequence step.
    * Work Group = swimlane; blank -> IT System. Work-group aliases from
      mapping/flow_rules.yaml merge lanes.
    * External/System work groups and the excluded task prefixes are skipped
      (diagram only; they stay in the task table).
    * A row whose explicit Seq is LOWER than an earlier row's Seq (the list
      goes backwards) is skipped and reported - the diagram never reorders
      the sequence.
    """
    rules = load_flow_rules()
    aliases = rules["lane_aliases"]
    excluded_wg = {str(x).lower() for x in rules["excluded_work_groups"]}
    prefixes = tuple(str(p).casefold() for p in rules["excluded_task_prefixes"])
    prepared: list[FlowTask] = []
    warnings: list[str] = []
    run_max = None

    for task in tasks:
        wg = _norm_ws(task.work_group)
        it_system = _norm_ws(task.it_system)
        task_name = _norm_ws(task.task_name)
        explicit = _clean(task.seq)
        eff = explicit or _clean(getattr(task, "flow_seq", None))

        if wg.lower() in excluded_wg:
            continue
        if any(task_name.casefold().startswith(p) for p in prefixes):
            continue
        it_lanes = {str(x).casefold() for x in rules["it_system_lanes"]}
        lane = it_system if it_system.casefold() in it_lanes else (wg or it_system)
        if not lane:
            continue
        if not eff:
            warnings.append(f"Task '{task_name}' has no Seq - left out of the flowchart.")
            continue
        if explicit:
            key = _natural_key(explicit)
            if run_max is not None and key < run_max:
                warnings.append(
                    f"Task '{task_name}' (Seq {explicit}, sheet row {getattr(task, 'source_row', '?')}) "
                    f"comes after a higher Seq in the Excel list - left out of the flowchart. "
                    f"Check its Seq number."
                )
                continue
            run_max = key
        prepared.append(FlowTask(
            seq=eff, task_name=task_name, it_system=it_system, work_group=wg,
            lane=_canon_lane(lane, aliases), condition=_norm_ws(task.task_condition),
            label=display_name(task_name), is_wa=is_work_activity(task_name),
        ))

    # Merge Work Groups that belong in one lane (only when 2+ of them are present).
    import dataclasses
    for rule in rules["lane_merges"]:
        members = {_norm_ws(x).casefold() for x in rule.get("lanes", [])}
        present = {t.lane.casefold() for t in prepared} & members
        if len(present) >= 2:
            title = _norm_ws(rule.get("title", ""))
            prepared = [dataclasses.replace(t, lane=title) if t.lane.casefold() in members else t
                        for t in prepared]

    # Stable sort: equal sequence numbers keep their Excel order.
    prepared.sort(key=lambda t: _natural_key(t.seq))
    return prepared, warnings


def prepare_flow_tasks(tasks: Iterable[Task]) -> list[FlowTask]:
    """Apply the flowchart rules and return sequence-sorted flow tasks."""
    return prepare_flow_tasks_ex(tasks)[0]


def _escape_xml_text(text: str) -> str:
    # lxml element text handles XML escaping for us; this helper is retained
    # for readable line-wrapping only.
    return text


def _wrap_text(text: str, max_chars: int) -> list[str]:
    words = _clean(text).split()
    if not words:
        return [""]
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def _append_textbox(shape, text: str, font_pt: int = 7, bold: bool = False):
    textbox = _v("textbox")
    textbox.set("inset", "2pt,2pt,2pt,2pt")

    txbx_content = OxmlElement("w:txbxContent")
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    p_pr.append(jc)
    p.append(p_pr)

    lines = _wrap_text(text, max(8, int(70 / max(font_pt, 1))))
    for idx, line in enumerate(lines):
        r = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_fonts.set(qn("w:hAnsi"), "Times New Roman")
        r_pr.append(r_fonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(font_pt * 2))
        r_pr.append(sz)
        if bold:
            r_pr.append(OxmlElement("w:b"))
        r.append(r_pr)
        if idx > 0:
            # w:br is only valid INSIDE a run (a bare w:br under w:p makes
            # Word report the file as corrupted).
            r.append(OxmlElement("w:br"))
        t = OxmlElement("w:t")
        t.text = line
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p.append(r)

    txbx_content.append(p)
    textbox.append(txbx_content)
    shape.append(textbox)


def _vml_shape(tag: str, shape_id: str, style: str, fill: str, stroke: str,
               text: str | None = None, rounded: bool = False,
               font_pt: int = 7, bold: bool = False):
    shape = _v(tag)
    shape.set("id", shape_id)
    shape.set("style", style)
    shape.set("fillcolor", f"#{fill}")
    shape.set("strokecolor", f"#{stroke}")
    if rounded:
        shape.set("arcsize", "0.18")
    if text is not None:
        _append_textbox(shape, text, font_pt=font_pt, bold=bold)
    return shape


def _vml_line(shape_id: str, x1: float, y1: float, x2: float, y2: float,
              arrow: bool = True):
    """Create one orthogonal connector segment.

    Keeping each segment as its own VML line preserves editability in Word.
    Arrowheads are placed only on the final segment of a multi-segment route.
    """
    line = _v("line")
    line.set("id", shape_id)
    line.set("from", f"{_pt(x1):.2f}pt,{_pt(y1):.2f}pt")
    line.set("to", f"{_pt(x2):.2f}pt,{_pt(y2):.2f}pt")
    line.set(
        "style",
        "position:absolute;z-index:2;stroke-width:1pt;"
    )
    line.set("strokecolor", f"#{CONNECTOR_COLOR}")
    if arrow:
        stroke = _v("stroke")
        stroke.set("endarrow", "block")
        stroke.set("weight", "1pt")
        line.append(stroke)
    return line


def _append_orthogonal_connector(pict, connector_id: str,
                                 start: tuple[float, float],
                                 end: tuple[float, float],
                                 prefer_vertical: bool = False):
    """Route a connector using horizontal/vertical segments rather than a
    single diagonal line. This makes cross-swimlane connections much easier
    to follow visually.
    """
    x1, y1 = start
    x2, y2 = end

    if abs(x1 - x2) < 1:
        pict.append(_vml_line(connector_id + "_v", x1, y1, x2, y2, arrow=True))
        return
    if abs(y1 - y2) < 1:
        pict.append(_vml_line(connector_id + "_h", x1, y1, x2, y2, arrow=True))
        return

    if prefer_vertical:
        mid_y = (y1 + y2) / 2
        pict.append(_vml_line(connector_id + "_v1", x1, y1, x1, mid_y, arrow=False))
        pict.append(_vml_line(connector_id + "_h", x1, mid_y, x2, mid_y, arrow=False))
        pict.append(_vml_line(connector_id + "_v2", x2, mid_y, x2, y2, arrow=True))
    else:
        mid_x = (x1 + x2) / 2
        pict.append(_vml_line(connector_id + "_h1", x1, y1, mid_x, y1, arrow=False))
        pict.append(_vml_line(connector_id + "_v", mid_x, y1, mid_x, y2, arrow=False))
        pict.append(_vml_line(connector_id + "_h2", mid_x, y2, x2, y2, arrow=True))


def _new_group_paragraph(doc: Document, group_id: str):
    """Create a paragraph whose VML shapes are independent siblings.

    The previous implementation placed every lane, task and connector inside
    one v:group. Word therefore treated the complete diagram as a single
    movable object. Here we deliberately keep only a w:pict container and add
    each VML shape directly to it, so shapes can be selected and moved
    independently in Word.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Reserve approximately the diagram height in the document flow. The
    # actual shapes are floating and positioned relative to this paragraph.
    p.paragraph_format.space_after = Inches(0.08)
    p.paragraph_format.line_spacing = Inches(GROUP_HEIGHT_IN)

    run = p.add_run()
    pict = OxmlElement("w:pict")
    run._r.append(pict)
    return p, pict

def _pt(value: float) -> float:
    """Convert the diagram's 100-units-per-inch coordinate to points."""
    return value * COORD_TO_PT

def _shape_style(x: float, y: float, w: float, h: float, z: int = 1) -> str:
    return (
        f"position:absolute;left:{_pt(x):.2f}pt;top:{_pt(y):.2f}pt;"
        f"width:{_pt(w):.2f}pt;height:{_pt(h):.2f}pt;z-index:{z};"
    )


def _add_panel(pict, flow_tasks: list[FlowTask], panel_no: int, total_panels: int):
    lanes: list[str] = []
    for task in flow_tasks:
        if task.lane not in lanes:
            lanes.append(task.lane)

    if not lanes:
        return

    lane_w = COORD_WIDTH / len(lanes)
    lane_body_h = COORD_HEIGHT - LANE_HEADER_H

    # Background and lane headers.
    for lane_idx, lane in enumerate(lanes):
        x = lane_idx * lane_w
        lane_shape = _vml_shape(
            "rect", f"lane_{panel_no}_{lane_idx}",
            _shape_style(x, 0, lane_w, COORD_HEIGHT, z=0),
            fill=LANE_FILL, stroke="9FC5E5",
        )
        pict.append(lane_shape)
        header = _vml_shape(
            "rect", f"lane_header_{panel_no}_{lane_idx}",
            _shape_style(x, 0, lane_w, LANE_HEADER_H, z=1),
            fill=LANE_HEADER_FILL, stroke="9FC5E5", text=lane,
            font_pt=7, bold=True,
        )
        pict.append(header)

    # Group tasks by their sequence number. In the source task list a
    # sequence can appear as e.g. ``10`` and ``10-ONBOARD CPE``. Both belong
    # to sequence 10, so the numeric prefix is used as the grouping key.
    # The FIRST task in each group is the task whose outgoing connector
    # continues to the next sequence.
    def sequence_group_key(seq: str) -> str:
        m = re.match(r"^(\d+(?:\.\d+)*)", _clean(seq))
        return m.group(1) if m else _clean(seq)

    sequence_groups: list[tuple[str, list[FlowTask]]] = []
    for task in flow_tasks:
        key = sequence_group_key(task.seq)
        if sequence_groups and sequence_groups[-1][0] == key:
            sequence_groups[-1][1].append(task)
        else:
            sequence_groups.append((key, [task]))

    available_h = lane_body_h - TOP_PAD - BOTTOM_PAD
    row_step = min(STEP_Y * 1.55, max(42, available_h / max(len(sequence_groups), 1)))
    task_positions: dict[int, tuple[float, float, float, float]] = {}
    task_index = {id(task): i for i, task in enumerate(flow_tasks)}
    group_positions: list[list[int]] = []

    for group_idx, (_seq, group) in enumerate(sequence_groups):
        lane_counts: dict[int, int] = {}
        for task in group:
            lane_idx = lanes.index(task.lane)
            lane_counts[lane_idx] = lane_counts.get(lane_idx, 0) + 1

        base_y = LANE_HEADER_H + TOP_PAD + group_idx * row_step
        # Keep duplicate tasks vertically close to each other, centered on the
        # sequence row. Single tasks remain centered on the row.
        spacing = min(42, max(30, row_step / max(len(group), 1)))
        group_height = (len(group) - 1) * spacing
        first_center_y = base_y - group_height / 2
        indices: list[int] = []

        for dup_idx, task in enumerate(group):
            idx = task_index[id(task)]
            indices.append(idx)
            lane_idx = lanes.index(task.lane)
            x_center = lane_idx * lane_w + lane_w / 2
            y_center = first_center_y + dup_idx * spacing
            common_w = min(FLOW_SHAPE_W, max(40, lane_w - 2 * LANE_PAD_X))
            box_w = common_w
            task_h = FLOW_SHAPE_H
            left = x_center - box_w / 2
            top = y_center - task_h / 2
            task_positions[idx] = (left, top, box_w, task_h)

            label = task.task_name
            font_pt = 6 if len(label) > 28 or box_w < 75 else 7
            shape_key = (_clean(task.seq).casefold(), _clean(task.task_name).casefold())
            semantic_shape = TASK_SHAPE_OVERRIDES.get(shape_key, "process")
            # Only standard, source-grounded process/decision task shapes are
            # supported here. Start/End remain synthetic terminators.
            vml_tag = "diamond" if semantic_shape == "decision" else "rect"
            shape = _vml_shape(
                vml_tag, f"task_{panel_no}_{idx}",
                _shape_style(left, top, box_w, task_h, z=3),
                fill=TASK_FILL, stroke=TASK_STROKE,
                text=label, font_pt=font_pt,
            )
            pict.append(shape)

        group_positions.append(indices)

    # Start marker points to the first eligible task.
    first = task_positions[group_positions[0][0]]
    first_task = flow_tasks[group_positions[0][0]]
    first_lane = lanes.index(first_task.lane)
    first_x = first_lane * lane_w + lane_w / 2
    first_y = LANE_HEADER_H + 15
    x1, y1, w1, h1 = first
    marker = _vml_shape(
        "oval", f"{'start' if panel_no == 1 else 'continue'}_{panel_no}",
        _shape_style(first_x - w1 / 2, first_y, w1, h1, z=4),
        fill=START_END_FILL, stroke=START_END_STROKE,
        text="Start" if panel_no == 1 else "Continue", font_pt=6,
    )
    pict.append(marker)
    _append_orthogonal_connector(
        pict, f"{'start' if panel_no == 1 else 'continue'}_line_{panel_no}",
        (first_x, first_y + h1), (x1 + w1 / 2, y1), prefer_vertical=True,
    )

    # Duplicate sequence numbers: connect the tasks vertically. This is a
    # deliberately visible up/down (vertical) relationship rather than a
    # diagonal line. The first task remains the primary task for the outgoing
    # connection to the next sequence.
    for indices in group_positions:
        if len(indices) > 1:
            for a, b in zip(indices, indices[1:]):
                xa, ya, wa, ha = task_positions[a]
                xb, yb, wb, hb = task_positions[b]
                _append_orthogonal_connector(
                    pict, f"duplicate_connector_{panel_no}_{a}_{b}",
                    (xa + wa / 2, ya + ha),
                    (xb + wb / 2, yb),
                    prefer_vertical=True,
                )

    # Connect each sequence's FIRST task to the FIRST task of the next
    # sequence. For different lanes, route with right-angle segments to avoid
    # ambiguous diagonal crossings.
    for group_idx in range(len(group_positions) - 1):
        a = group_positions[group_idx][0]
        b = group_positions[group_idx + 1][0]
        xa, ya, wa, ha = task_positions[a]
        xb, yb, wb, hb = task_positions[b]
        same_lane = lanes.index(flow_tasks[a].lane) == lanes.index(flow_tasks[b].lane)
        _append_orthogonal_connector(
            pict, f"sequence_connector_{panel_no}_{group_idx}",
            (xa + wa / 2, ya + ha),
            (xb + wb / 2, yb),
            prefer_vertical=not same_lane,
        )

    # End marker follows the LAST sequence's FIRST task, because that is the
    # designated continuation task for each sequence.
    last = task_positions[group_positions[-1][0]]
    end_task = flow_tasks[group_positions[-1][0]]
    end_lane = lanes.index(end_task.lane)
    end_x = end_lane * lane_w + lane_w / 2
    end_y = min(COORD_HEIGHT - 30, last[1] + last[3] + 20)
    if panel_no == total_panels:
        end = _vml_shape(
            "oval", f"end_{panel_no}",
            _shape_style(end_x - last[2] / 2, end_y, last[2], last[3], z=4),
            fill=START_END_FILL, stroke=START_END_STROKE,
            text="End", font_pt=6,
        )
        pict.append(end)
        _append_orthogonal_connector(
            pict, f"end_line_{panel_no}",
            (last[0] + last[2] / 2, last[1] + last[3]),
            (end_x, end_y), prefer_vertical=True,
        )

    if total_panels > 1:
        label = _vml_shape(
            "rect", f"panel_label_{panel_no}",
            _shape_style(4, COORD_HEIGHT - 18, 110, 14, z=5),
            fill=LANE_FILL, stroke=LANE_FILL,
            text=f"Flow continued - {panel_no}/{total_panels}", font_pt=6,
        )
        pict.append(label)


def add_task_flowchart(doc: Document, tasks: Iterable[Task]) -> dict:
    """Insert one or more editable swimlane flowchart panels.

    Returns summary metadata useful for logging/testing.
    """
    flow_tasks = prepare_flow_tasks(tasks)
    if not flow_tasks:
        p = doc.add_paragraph()
        r = p.add_run("[No eligible tasks for flowchart: External/System tasks were excluded, or no swimlane could be assigned.]")
        r.italic = True
        return {"eligible_tasks": 0, "panels": 0, "lanes": []}

    lanes = []
    for task in flow_tasks:
        if task.lane not in lanes:
            lanes.append(task.lane)

    panels = [
        flow_tasks[i:i + MAX_TASKS_PER_PANEL]
        for i in range(0, len(flow_tasks), MAX_TASKS_PER_PANEL)
    ]

    for panel_idx, panel_tasks in enumerate(panels, start=1):
        if panel_idx > 1:
            doc.add_page_break()
            p = doc.add_paragraph()
            r = p.add_run("Task Flow Diagram (continued)")
            r.bold = True
            r.font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p, pict = _new_group_paragraph(doc, f"task_flow_{id(tasks)}_{panel_idx}")
        _add_panel(pict, panel_tasks, panel_idx, len(panels))
        # Panel numbers restart at 1 for every order type, so make every VML
        # shape id unique across the whole document.
        uid = next(_PANEL_UID)
        for el in pict:
            if el.get("id"):
                el.set("id", f"{el.get('id')}_d{uid}")
        p.paragraph_format.space_after = Pt(4)

    return {
        "editable_independent_shapes": True,
        "eligible_tasks": len(flow_tasks),
        "panels": len(panels),
        "lanes": lanes,
    }
