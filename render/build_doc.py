"""
Builds the Service Delivery Guide .docx from an ExtractionResult.

Fully product-agnostic: every heading, title, and label is built from
`result.product_name` - nothing here is hardcoded to a specific product.

Document structure:
  1. Front page (no header) - product name (36pt) / "SERVICE DELIVERY
     GUIDE" (18pt)
  2. Document Control Sheet pages (contact block, sign-off tables,
     revision history, summary of changes)
  3. Table of Contents (auto-updating field)
  4. Section 1 onward, as before - now with a running header on every
     page except the front page, Times New Roman throughout, and
     per-task explanation pages after each order type's task table.
"""
from __future__ import annotations
import os
from copy import deepcopy
from lxml import etree
from datetime import date
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from models.schemas import ExtractionResult, OrderTypeSection, Task
from render import components as C
from render import template_utils
from render.task_flow import add_task_flowchart  # legacy editable-VML version (no longer used)
from render.task_flow_image import add_task_flowchart_image
import tempfile

_FLOW_IMG_DIR = tempfile.mkdtemp(prefix="sdg_flow_")

ASH_FILL = C.ASH_FILL
BLUE = C.HEADER_FILL
LIGHT_BLUE = "A9C6E8"  # lighter tint used for front-page lines and header background

# --- Task table: only these 4 columns per spec (Duration / Cancellation WO /
# Task Condition are still in the data model and available if needed later,
# just not shown in this table) ---
TASK_TABLE_HEADERS = ["Task Seq.", "Task Name", "IT System", "Work Group"]
TASK_TABLE_WIDTHS_IN = [0.9, 2.6, 1.2, 1.8]

# Task explanation pages are skipped for tasks whose Work Group is blank,
# exactly "SYSTEM" or "EXTERNAL", or contains "auto" anywhere
# (case-insensitive) - this catches "AUTO", "AUTO<Hidden task>",
# "AUTO (hidden)", etc. Per spec, those are automated / externally
# handled / not owned by an internal team and don't need a procedural
# write-up.
EXCLUDED_WORK_GROUPS_EXACT = {"", "system", "external"}

DETAIL_LABELS = ["Task Owner", "Description", "Start Event", "End Event", "KPI", "OLA"]

ASSETS_DIR = os.path.join(os.path.dirname(__file__), "..", "assets")
DEFAULT_LOGO_PATH = os.path.join(ASSETS_DIR, "placeholder_logo.png")


def _task_rows(section: OrderTypeSection) -> list[list[str]]:
    return [[t.seq, t.task_name, t.it_system, t.work_group] for t in section.tasks]


def _qualifies_for_explanation(task: Task) -> bool:
    wg = (task.work_group or "").strip().lower()
    if wg in EXCLUDED_WORK_GROUPS_EXACT:
        return False
    if "auto" in wg:
        return False
    return True


def _set_base_styles(doc: Document):
    """Times New Roman throughout: body 12pt, headings 14pt bold, all black."""
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "Times New Roman")

    for lvl in range(1, 5):
        try:
            h = doc.styles[f"Heading {lvl}"]
            h.font.name = "Times New Roman"
            h.font.size = Pt(14)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        except KeyError:
            pass

    try:
        lb = doc.styles["List Bullet"]
        lb.font.name = "Times New Roman"
        lb.font.size = Pt(12)
        lb.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    except KeyError:
        pass


def _force_times_new_roman(doc: Document):
    """Force Times New Roman across the supplied template and generated content.

    This includes normal/heading/list styles, ordinary runs, table text,
    header/footer text and text embedded in Word text boxes.  The template
    remains the layout master; this function only enforces the requested
    document-wide font family.
    """
    font_name = "Times New Roman"
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def set_rfonts(rpr):
        if rpr is None:
            return
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font_name)

    # Styles
    for style in doc.styles:
        try:
            style.font.name = font_name
            rpr = style._element.get_or_add_rPr()
            set_rfonts(rpr)
        except Exception:
            pass

    # Body + headers/footers, including text boxes/VML paragraphs.
    roots = [doc._body._element]
    for section in doc.sections:
        roots.extend([
            section.header._element,
            section.first_page_header._element,
            section.footer._element,
            section.first_page_footer._element,
        ])

    for root in roots:
        for r in root.xpath(".//w:r"):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                r.insert(0, rpr)
            set_rfonts(rpr)


def _build_running_header(
    section, product: str, doc_no: str, issue_no: str,
    issue_date: str, revision_no: str, revision_date: str, logo_path: str | None,
):
    """Header shown on every page except the first page of the document.
    Layout: [logo] | [product name (blue) / Document No, centered] |
    [Issue No., Date of Issue, Revision No., Date of Revision, Page X of Y
    - left-aligned]
    """
    section.different_first_page_header_footer = True
    header = section.header
    header.is_linked_to_previous = False

    table = header.add_table(rows=1, cols=3, width=Inches(6.5))
    table.autofit = False
    widths = [Inches(1.4), Inches(3.2), Inches(1.9)]
    for cell, w in zip(table.rows[0].cells, widths):
        cell.width = w

    cell_logo, cell_center, cell_right = table.rows[0].cells
    for cell in (cell_logo, cell_center, cell_right):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Light-blue background behind the product-name/document-number cell,
    # same tint as the front-page lines
    C.set_cell_shading(cell_center, LIGHT_BLUE)

    # Left: logo
    p = cell_logo.paragraphs[0]
    run = p.add_run()
    logo = logo_path or DEFAULT_LOGO_PATH
    if os.path.exists(logo):
        try:
            run.add_picture(logo, width=Inches(1.1))
        except Exception as e:
            print(f"      WARNING: could not insert logo image '{logo}' into "
                  f"the header ({e}). Check the file is a valid PNG/JPEG.")
    else:
        print(f"      WARNING: logo path '{logo}' does not exist - header "
              f"will have no logo.")

    # Center: product name (black text, blue background) + document number
    p1 = cell_center.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(product)
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.name = "Times New Roman"
    r1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p2 = cell_center.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"Document No: {doc_no}")
    r2.font.size = Pt(9)
    r2.font.name = "Times New Roman"
    r2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Right: issue/revision metadata + page number, stacked, left-aligned
    right_lines = [
        f"Issue No.: {issue_no}",
        f"Date of Issue: {issue_date}",
        f"Revision No.: {revision_no}",
        f"Date of Revision: {revision_date}",
    ]
    p = cell_right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(right_lines[0])
    r.font.size = Pt(8)
    r.font.name = "Times New Roman"
    for line in right_lines[1:]:
        np = cell_right.add_paragraph()
        np.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = np.add_run(line)
        r.font.size = Pt(8)
        r.font.name = "Times New Roman"
    pg = cell_right.add_paragraph()
    pg.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = pg.add_run("Page ")
    r.font.size = Pt(8)
    r.font.name = "Times New Roman"
    C.add_field(pg, "PAGE")
    r = pg.add_run(" of ")
    r.font.size = Pt(8)
    r.font.name = "Times New Roman"
    C.add_field(pg, "NUMPAGES")

    # strip the default header table borders for a clean look
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)


def _add_document_control_pages(doc: Document, product: str, meta: dict):
    """Document Control Sheet: contact block, sign-off tables, revision
    history, summary of changes. Placed before the Table of Contents."""
    doc.add_heading("Document Control Sheet", level=1)

    doc.add_heading("Contact for Inquiries & Proposed Changes", level=2)
    C.add_label_value_table(doc, [
        ("Name", meta.get("contact_name", "")),
        ("Designation", meta.get("contact_designation", "DGM/BPM")),
        ("Telephone", meta.get("contact_telephone", "")),
        ("Mobile", meta.get("contact_mobile", "")),
        ("Extension", meta.get("contact_extension", "")),
        ("E-mail", meta.get("contact_email", "")),
    ])
    doc.add_paragraph()

    C.add_signoff_table(doc, "Prepared By")
    doc.add_paragraph()
    C.add_signoff_table(doc, "Checked By")
    doc.add_paragraph()
    C.add_signoff_table(doc, "Approved By")
    doc.add_paragraph()

    doc.add_heading("Revision History", level=2)
    C.add_table(doc, ["Issue Number", "Issued Date", "Issued By"], [["", "", ""]],
                col_widths_in=[1.8, 1.8, 2.9])
    doc.add_paragraph()

    doc.add_heading("Summary of Changes", level=2)
    C.add_table(doc, ["Document", "Section"], [["", ""]], col_widths_in=[3.25, 3.25])

    C.add_page_break(doc)


def _add_task_explanation_page(doc: Document, task: Task):
    """Add the editable task-detail table for one qualifying task.

    Task Owner is populated from the task's Work Group and Description is
    populated from sheet 6.4 when a matching description was found.  The
    remaining fields intentionally stay blank for manual completion.
    """
    C.add_page_break(doc)

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"

    header_cell = C.merge_cells_horizontal(table, 0, 0, 1)
    header_cell.paragraphs[0].text = ""
    run = header_cell.paragraphs[0].add_run(
        f"{task.seq or ''}   {task.task_name}".strip()
    )
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = "Times New Roman"
    header_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    C.set_cell_shading(header_cell, BLUE)

    values = {
        "Task Owner": task.work_group or "",
        "Description": task.description or "",
        "Start Event": "",
        "End Event": "",
        "KPI": "",
        "OLA": "",
    }

    for i, label in enumerate(DETAIL_LABELS, start=1):
        cells = table.rows[i].cells
        label_p = cells[0].paragraphs[0]
        label_p.text = ""
        label_run = label_p.add_run(label)
        label_run.bold = True
        label_run.font.name = "Times New Roman"

        value_p = cells[1].paragraphs[0]
        value_p.text = str(values[label])
        for value_run in value_p.runs:
            value_run.font.name = "Times New Roman"

    doc.add_paragraph()
    for line in ("Business Rule:", "Work Instructions:", "Required Information & Sources:"):
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.bold = True
        r.font.name = "Times New Roman"




def _add_dynamic_content(
    doc: Document,
    result: ExtractionResult,
    product: str,
    contact_meta: dict,
    template_mode: bool = False,
    include_task_flowchart: bool = True,
):
    """Append only the data-driven part of the guide.

    The template owns the cover page, running header/footer, document-control
    layout, contents-page layout, and the Introduction anchor. This function
    owns only content that changes with the Excel/YAML input.
    """
    def heading(title: str, level: int, number: str | None = None):
        # Generated numbering is explicit and deterministic. The supplied
        # template's Heading 1-3 styles contain automatic numbering, so the
        # renderer uses numbering-free visual copies of those styles. This
        # prevents Word from continuing the template's sample numbering and
        # producing values such as 5, 12.0, etc.
        if number:
            prefix = f"{number}." if level == 1 else number
            text = f"{prefix} {title}"
        else:
            text = title
        return C.add_heading(doc, text, level)

    # --- Section 1 content: the template already contains the visible
    # "01. Introduction" heading. Its body is populated here.
    introduction_content = result.ai_introduction or result.introduction_text
    if introduction_content:
        for para in introduction_content:
            doc.add_paragraph(para)
    else:
        C.add_stub(doc, "1. Product Description")

    heading("Product-related Business Rules", 2, "1.1")
    if result.business_rule_blocks:
        # Render each requirement as text, then render only the table(s) that
        # belong to that requirement. The overall Special Requirements sheet
        # is never flattened into one large Word table.
        for block in result.business_rule_blocks:
            p = doc.add_paragraph(block.text)
            p.paragraph_format.space_after = Pt(4)
            for table_data in block.tables:
                C.add_excel_table(doc, table_data.headers, table_data.rows, group_headers=table_data.group_headers)
                doc.add_paragraph()
    elif result.business_rules:
        for rule in result.business_rules:
            doc.add_paragraph(rule, style="List Paragraph")
    else:
        C.add_stub(doc, "3. Special Requirements ('Requirement Description' column)")
    C.add_page_break(doc)

    # --- Section 2: Product Hierarchy ---
    heading("Product Hierarchy", 1, "2")
    # The narrative/list of product components is intentionally omitted.
    # Product Hierarchy now proceeds directly to attributes and values.

    if result.hierarchy_images:
        for img_path in result.hierarchy_images:
            if os.path.exists(img_path):
                doc.add_picture(img_path, width=Inches(5.5))

    heading("Product Attributes with Definitions", 2, "2.1")
    if result.hierarchy:
        for i, node in enumerate(result.hierarchy, start=1):
            heading(node.name, 3, f"2.1.{i}")

            heading("Attribute Definitions", 4, f"2.1.{i}.1")
            if node.attributes:
                attr_rows = [[a.entity_name, a.mandatory, a.comments] for a in node.attributes]
                C.add_table(
                    doc, ["Entity Name", "Mandatory?", "Comments"], attr_rows,
                    col_widths_in=[2.2, 1.2, 3.4],
                )
            else:
                C.add_stub(
                    doc,
                    f"2. Product Hierarchy (no 'Attribute' rows found under {node.name})",
                )

            heading("Attribute List of Values", 4, f"2.1.{i}.2")
            attrs_with_values = [
                (a.entity_name, a.values) for a in node.attributes if a.values
            ]
            if attrs_with_values:
                C.add_attribute_values_table(doc, attrs_with_values)
            else:
                C.add_stub(
                    doc,
                    f"2. Product Hierarchy (no 'Value' rows found for any "
                    f"attribute under {node.name})",
                )
    else:
        C.add_stub(doc, "2. Product Hierarchy")
    C.add_page_break(doc)

    # --- Section 3: Order Types with Descriptions ---
    heading("Order Types with Descriptions", 1, "3")
    doc.add_paragraph(
        f"The following order types are required in the end-to-end "
        f"environment to be configured for fulfilling the business "
        f"requirements of the {product} product."
    )
    if result.order_scenarios:
        scenario_rows = [
            [s.business_scenario, s.crm_order_sub_type, s.crm_order_condition, s.oss_service_order]
            for s in result.order_scenarios
        ]
    else:
        scenario_rows = [["[pending]", "[pending]", "[pending]", "[pending]"] for _ in range(5)]

    C.add_table(
        doc,
        ["Business Scenario", "CRM Order Sub Type", "CRM Order Condition", "OSS Service Order"],
        scenario_rows,
        col_widths_in=[2.0, 1.6, 1.6, 1.6],
    )
    if not result.order_scenarios:
        p = doc.add_paragraph()
        run = p.add_run(
            "[No rows with an Order Type were found in sheet "
            "5.1 MDM_Order Scenarios.]"
        )
        run.italic = True
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    C.add_page_break(doc)

    # --- Sections 4..N: Workflow Diagram per order type ---
    for i, sec in enumerate(result.order_types):
        n = i + 4
        heading(f"Workflow Diagram - {sec.order_type} Order Type", 1, str(n))
        heading(product, 2, f"{n}.1")

        # The flowchart belongs directly under 4.1 / 5.1 / 6.1 ... .
        # The detailed task-flow heading is deliberately placed AFTER the
        # diagram so it does not become the diagram's parent topic.
        p = doc.add_paragraph()
        run = p.add_run(
            f"The {sec.order_type} order type consists of {sec.task_count} "
            f"tasks, executed in the sequence below."
        )
        run.italic = True

        if include_task_flowchart:
            flow_title = doc.add_paragraph()
            flow_run = flow_title.add_run("Task Flow Diagram")
            flow_run.bold = True
            flow_run.font.size = Pt(10)
            summary = add_task_flowchart_image(doc, sec.tasks, sec.order_type, _FLOW_IMG_DIR, product)
            for w in summary.get("warnings", []):
                result.warnings.append(f"[{sec.order_type}] Flowchart: {w}")
            if summary["eligible_tasks"] == 0:
                result.warnings.append(
                    f"[{sec.order_type}] No eligible tasks were available for the task-flow diagram."
                )

        heading(f"Task Flow - {product}", 3, f"{n}.1.1")

        table_title = doc.add_paragraph()
        table_run = table_title.add_run("Task Details")
        table_run.bold = True
        table_run.font.size = Pt(10)

        C.add_table(
            doc,
            TASK_TABLE_HEADERS,
            _task_rows(sec),
            col_widths_in=TASK_TABLE_WIDTHS_IN,
        )

        for task in sec.tasks:
            if _qualifies_for_explanation(task):
                _add_task_explanation_page(doc, task)

        C.add_page_break(doc)

    # --- Final section: Fault Handling Process ---
    fault_n = len(result.order_types) + 4
    heading("Fault Handling Process", 1, str(fault_n))
    C.add_stub(doc, "Source sheet not yet identified — confirm with document owner")


def _remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _paragraph_text_xml(p):
    return "".join(p.xpath(".//w:t/text()"))


def _append_value_to_xml_paragraph(p, value: str):
    """Append a value to an existing formatted XML paragraph without
    disturbing the template's label/tabs/formatting."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    runs = p.xpath("./w:r")
    last_run = None
    for r in reversed(runs):
        if r.xpath(".//w:t"):
            last_run = r
            break
    new_r = OxmlElement("w:r")
    if last_run is not None:
        rpr = last_run.find(qn("w:rPr"))
        if rpr is not None:
            new_r.append(deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = str(value or "")
    new_r.append(t)
    p.append(new_r)


def _replace_last_text_in_xml_paragraph(p, value: str):
    """Replace the final text-bearing run in a paragraph, preserving its
    existing run formatting."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    texts = p.xpath(".//w:t")
    if not texts:
        _append_value_to_xml_paragraph(p, value)
        return
    texts[-1].text = str(value or "")


def _iter_all_paragraphs(part_element):
    """Yield all w:p elements, including paragraphs inside Word text boxes."""
    return part_element.xpath(".//w:p")


def _populate_template_metadata(
    doc: Document,
    product: str,
    doc_no: str,
    issue_no: str,
    issue_date: str,
    revision_no: str,
    revision_date: str,
    issued_by: str,
    contact_meta: dict,
):
    """Fill the supplied Word template in-place.

    The uploaded template is intentionally treated as a visual master rather
    than rebuilt. Its cover and header are Word shapes/text boxes, so the
    metadata is populated at XML level to preserve the exact layout.
    """
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    # 1) Body/cover text boxes and normal paragraphs.
    body = doc._body._element
    for p in _iter_all_paragraphs(body):
        text = _paragraph_text_xml(p).strip()

        if text == "Title":
            _replace_last_text_in_xml_paragraph(p, product)
        elif text == "Document Number: BPMS XXX XXX":
            # Replace only the value, preserving the label and formatting.
            texts = p.xpath(".//w:t")
            if texts:
                texts[-1].text = str(doc_no)
        elif text in {"Issue No:", "Issue No  :", "Issue No  :"}:
            _append_value_to_xml_paragraph(p, issue_no)
        elif text.startswith("Issue Date:") or text.replace(" ", "").startswith("IssueDate:"):
            _append_value_to_xml_paragraph(p, issue_date)
        elif text.startswith("Issued By:") or text.replace(" ", "").startswith("IssuedBy:"):
            _append_value_to_xml_paragraph(p, issued_by)
        elif text == "Name:":
            _append_value_to_xml_paragraph(p, contact_meta.get("contact_name", ""))
        elif text == "Designation:":
            _append_value_to_xml_paragraph(p, contact_meta.get("contact_designation", ""))
        elif text.startswith("Telephone:"):
            _append_value_to_xml_paragraph(p, contact_meta.get("contact_telephone", ""))
        elif text.startswith("Mobile:"):
            _append_value_to_xml_paragraph(p, contact_meta.get("contact_mobile", ""))
        elif text.startswith("Extension:"):
            _append_value_to_xml_paragraph(p, contact_meta.get("contact_extension", ""))
        elif text.startswith("E-mail:"):
            _append_value_to_xml_paragraph(p, contact_meta.get("contact_email", ""))

    # 2) Regular running header (the template uses drawing/text-box shapes).
    for section in doc.sections:
        for header in (section.header, section.first_page_header):
            root = header._element
            for p in _iter_all_paragraphs(root):
                text = _paragraph_text_xml(p).strip()
                if not text:
                    continue

                if text.endswith("Title") and "Document No" not in text:
                    _replace_last_text_in_xml_paragraph(p, product)
                elif text == "Document No:":
                    _append_value_to_xml_paragraph(p, doc_no)
                elif text.startswith("Issue No.:"):
                    _replace_last_text_in_xml_paragraph(p, f": {issue_no}")
                elif text.startswith("Date of Issue"):
                    _replace_last_text_in_xml_paragraph(p, f": {issue_date}")
                elif text.startswith("Revision No"):
                    _replace_last_text_in_xml_paragraph(p, f": {revision_no}")
                elif text.startswith("Date of Revision"):
                    _replace_last_text_in_xml_paragraph(p, f": {revision_date}")

    # 3) Revision history first editable row in the template.
    # Template table order is stable: contact/sign-off tables first,
    # revision history is the final body table.
    if doc.tables:
        revision_table = doc.tables[-1]
        if len(revision_table.rows) >= 2 and len(revision_table.columns) >= 5:
            values = [
                issue_no,
                issue_date,
                issued_by,
                "",
                "",
            ]
            for idx, value in enumerate(values):
                revision_table.cell(1, idx).text = str(value)

    # 4) Ask Word to refresh fields such as PAGE/NUMPAGES and the TOC.
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _prepare_template_contents_page(doc: Document):
    """Replace the template's sample TOC entries with a live Word TOC field,
    while keeping the template's CONTENT heading/page layout."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = doc._body._element

    target_sdt = None
    for sdt in body.xpath("./w:sdt"):
        text = "".join(node.text or "" for node in sdt.iter(qn("w:t"))).strip()
        if "CONTENT" in text.upper():
            target_sdt = sdt
            break

    if target_sdt is None:
        return False

    content = target_sdt.find(qn("w:sdtContent"))
    if content is None:
        return False

    content_paras = [child for child in content if child.tag == qn("w:p")]
    content_heading = None
    for p in content_paras:
        if _paragraph_text_xml(p).strip().upper() == "CONTENT":
            content_heading = p
            break

    if content_heading is None:
        return False

    # Keep only the heading paragraph and replace everything after it with
    # a live TOC field.
    for child in list(content):
        if child is not content_heading:
            _remove_element(child)

    toc_doc = Document()
    C.add_toc(toc_doc)
    toc_p = toc_doc._body._element.find(qn("w:p"))
    content.append(deepcopy(toc_p))
    return True


def _insert_generated_nodes_after_anchor(doc: Document, anchor_text: str, generated_start_index: int):
    """Move newly generated body elements from the end of the template to
    immediately after the requested anchor paragraph."""
    body = doc._body._element
    children = list(body)
    generated = [
        child for child in children[generated_start_index:]
        if etree.QName(child).localname != "sectPr"
    ]

    anchor = None
    for child in children:
        if etree.QName(child).localname != "p":
            continue
        if _paragraph_text_xml(child).strip() == anchor_text:
            anchor = child
            break

    if anchor is None:
        return False

    # Remove generated nodes from their current positions, then insert in
    # original order immediately after the anchor.
    for child in generated:
        body.remove(child)

    insert_after = anchor
    for child in generated:
        insert_after.addnext(child)
        insert_after = child

    return True


def _remove_blank_paragraphs_after_anchor(doc: Document, anchor_text: str):
    body = doc._body._element
    children = list(body)
    anchor_index = None
    for i, child in enumerate(children):
        if etree.QName(child).localname == "p" and _paragraph_text_xml(child).strip() == anchor_text:
            anchor_index = i
            break
    if anchor_index is None:
        return

    # Remove only empty paragraphs immediately following the anchor, leaving
    # sectPr untouched. This prevents the template's placeholder whitespace
    # from pushing the generated Introduction content down the page.
    for child in list(children[anchor_index + 1:]):
        if etree.QName(child).localname != "p":
            break
        if _paragraph_text_xml(child).strip():
            break
        _remove_element(child)


def _build_from_scratch(
    result: ExtractionResult,
    output_path: str,
    product: str,
    doc_version: str | None,
    doc_no: str,
    issue_no: str,
    issue_date: str,
    revision_no: str,
    revision_date: str,
    issued_by: str,
    logo_path: str | None,
    contact_meta: dict,
    include_task_flowchart: bool = True,
):
    """Original code-generated layout, retained as a fallback."""
    doc = Document()
    _set_base_styles(doc)
    _build_running_header(
        doc.sections[0], product, doc_no, issue_no, issue_date,
        revision_no, revision_date, logo_path,
    )

    # --- Front page ---
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = logo_p.add_run()
    logo = logo_path or DEFAULT_LOGO_PATH
    if os.path.exists(logo):
        try:
            run.add_picture(logo, width=Inches(1.4))
        except Exception as e:
            print(f"      WARNING: could not insert logo image '{logo}' into the front page ({e}).")
    else:
        print(f"      WARNING: logo path '{logo}' does not exist - front page will have no logo.")

    C.add_horizontal_rule(doc, color_hex=LIGHT_BLUE, size_pt=3)

    def _front_page_paragraph(space_before_pt=0):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(1.15)
        p.paragraph_format.space_before = Pt(space_before_pt)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "12")
        left.set(qn("w:color"), LIGHT_BLUE)
        p_bdr.append(left)
        p_pr.append(p_bdr)
        return p

    def _add_run(p, text, size_pt, bold=False):
        run = p.add_run(text)
        run.font.size = Pt(size_pt)
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        run.bold = bold
        return run

    p = _front_page_paragraph(space_before_pt=24)
    _add_run(p, product, 36, bold=True)
    p = _front_page_paragraph()
    _add_run(p, "SERVICE DELIVERY GUIDE", 18)
    for _ in range(6):
        _front_page_paragraph()

    def _cover_field(label: str, value: str):
        p = _front_page_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(2.8))
        _add_run(p, f"{label}\t: {value}", 12)

    _cover_field("Document Number", doc_no)
    _cover_field("Prepared By", "")
    _cover_field("Reviewed By", "")
    _cover_field("Approved By", "")
    _front_page_paragraph()
    _front_page_paragraph()
    _cover_field("Issue No", issue_no)
    _cover_field("Issue Date", issue_date)
    _cover_field("Issued By", issued_by)
    for _ in range(14):
        _front_page_paragraph()
    C.add_page_break(doc)

    _add_document_control_pages(doc, product, contact_meta)
    doc.add_heading("Contents", level=1)
    C.add_toc(doc)
    C.add_page_break(doc)

    # Generate data-driven content at the end.
    start = len(doc._body._element)
    _add_dynamic_content(doc, result, product, contact_meta, template_mode=False, include_task_flowchart=include_task_flowchart)
    doc.save(output_path)
    return output_path


def build(
    result: ExtractionResult,
    output_path: str,
    doc_version: str | None = None,
    doc_no: str = "[TBD]",
    issue_no: str = "[TBD]",
    issue_date: str | None = None,
    revision_no: str = "[TBD]",
    revision_date: str = "[TBD]",
    issued_by: str = "BPM Section",
    logo_path: str | None = None,
    contact_name: str = "",
    contact_designation: str = "DGM/BPM",
    contact_telephone: str = "",
    contact_mobile: str = "",
    contact_extension: str = "",
    contact_email: str = "",
    template_path: str | None = None,
    include_task_flowchart: bool = True,
) -> str:
    product = result.product_name
    issue_date = issue_date or date.today().isoformat()

    contact_meta = dict(
        contact_name=contact_name,
        contact_designation=contact_designation,
        contact_telephone=contact_telephone,
        contact_mobile=contact_mobile,
        contact_extension=contact_extension,
        contact_email=contact_email,
    )

    if template_path:
        doc = Document(template_path)

        # IMPORTANT: do not call _set_base_styles() here. The Word template is
        # the source of truth for fonts, heading colours, spacing, page setup,
        # cover design, header/footer and table styles.

        # Support optional {{TOKEN}} placeholders as an additional,
        # template-author-friendly mechanism. The real supplied template does
        # not require tokens because its existing labels/shapes are populated
        # by _populate_template_metadata().
        template_utils.replace_placeholders(doc, {
            "{{PRODUCT_NAME}}": product,
            "{{DOCUMENT_NO}}": doc_no,
            "{{ISSUE_NO}}": issue_no,
            "{{ISSUE_DATE}}": issue_date,
            "{{REVISION_NO}}": revision_no,
            "{{REVISION_DATE}}": revision_date,
            "{{ISSUED_BY}}": issued_by,
        })

        _populate_template_metadata(
            doc, product, doc_no, issue_no, issue_date,
            revision_no, revision_date, issued_by, contact_meta,
        )

        if not _prepare_template_contents_page(doc):
            result.warnings.append(
                "Template CONTENT page could not be located; the template's "
                "existing contents page was left unchanged."
            )

        # Remove placeholder whitespace after the Introduction anchor.
        _remove_blank_paragraphs_after_anchor(doc, "Introduction")

        # Generate all data-driven content at the end so that python-docx
        # relationships (especially inserted images) belong to the real
        # target document, then move those nodes behind the template anchor.
        body = doc._body._element
        generated_start_index = len(body)
        _add_dynamic_content(doc, result, product, contact_meta, template_mode=True, include_task_flowchart=include_task_flowchart)

        if not _insert_generated_nodes_after_anchor(
            doc, "Introduction", generated_start_index
        ):
            result.warnings.append(
                "Template anchor 'Introduction' was not found. Generated "
                "content was left at the end of the template."
            )

        # User requirement: Times New Roman throughout the final document,
        # including the supplied template and all newly generated content.
        _force_times_new_roman(doc)

        doc.save(output_path)
        return output_path

    return _build_from_scratch(
        result, output_path, product, doc_version,
        doc_no, issue_no, issue_date, revision_no, revision_date,
        issued_by, logo_path, contact_meta, include_task_flowchart,
    )
