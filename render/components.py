"""
Reusable python-docx building blocks. build_doc.py composes these;
nothing here should know about "order types" or any domain-specific
concept - keep it generic so it's reusable for sections 1-3 and 13 too.
"""
from __future__ import annotations
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from copy import deepcopy
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEADER_FILL = "1F4E78"
ALT_ROW_FILL = "F2F2F2"
ASH_FILL = "A6A6A6"
STUB_COLOR = RGBColor(0x99, 0x99, 0x99)
FONT_NAME = "Times New Roman"


def add_horizontal_rule(doc: Document, color_hex: str = HEADER_FILL, size_pt: int = 4):
    """Full-width horizontal rule (paragraph bottom border), spanning the
    page's text area between margins."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_pt * 8))  # eighths of a point
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return p


def _ensure_plain_heading_style(doc: Document, source_name: str) -> str:
    """Create a numbering-free copy of a template heading style.

    The supplied Word template uses automatic multilevel numbering on
    Heading 1/2/3. Generated sections need deterministic numbers such as
    4, 4.1 and 4.1.1, so the renderer writes the numbers explicitly and
    uses visual copies of the template heading styles without numPr.
    """
    target_name = f"{source_name} Plain"
    try:
        doc.styles[target_name]
        return target_name
    except KeyError:
        pass

    source = doc.styles[source_name]
    target = doc.styles.add_style(target_name, WD_STYLE_TYPE.PARAGRAPH)
    target.base_style = doc.styles["Normal"]

    # Copy the source paragraph properties except automatic numbering.
    source_ppr = source._element.pPr
    if source_ppr is not None:
        target_ppr = target._element.get_or_add_pPr()
        for child in source_ppr:
            if child.tag == qn("w:numPr"):
                continue
            target_ppr.append(deepcopy(child))

    # Copy the source run/font properties so the template remains the
    # source of truth for heading appearance.
    source_rpr = source._element.rPr
    if source_rpr is not None:
        target_rpr = target._element.get_or_add_rPr()
        for child in source_rpr:
            target_rpr.append(deepcopy(child))

    return target_name


def add_heading(doc: Document, text: str, level: int):
    """Add a heading using the supplied template's visual formatting.

    Levels 1-3 use numbering-free copies of the template's Heading 1-3
    styles. The section number is supplied explicitly by the renderer.
    Level 4+ is intentionally bold Normal text because the supplied
    template has no suitable Heading 4/5 style.
    """
    if level <= 3:
        style_name = _ensure_plain_heading_style(doc, f"Heading {level}")
        return doc.add_paragraph(text, style=style_name)

    paragraph = doc.add_paragraph(text, style="Normal")
    for run in paragraph.runs:
        run.bold = True
    return paragraph


def add_stub(doc: Document, missing_source: str):
    """A visibly-marked placeholder for content whose source isn't wired up yet."""
    p = doc.add_paragraph()
    run = p.add_run(f"[Content pending — source not yet provided: {missing_source}]")
    run.italic = True
    run.font.color.rgb = STUB_COLOR
    run.font.name = FONT_NAME
    return p


def add_page_break(doc: Document):
    doc.add_page_break()


def set_cell_shading(cell, hex_color: str):
    """Public wrapper - set a table cell's background fill color."""
    _set_cell_shading(cell, hex_color)


def _set_cell_shading(cell, hex_color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths_in: list[float] | None = None,
              header_fill_overrides: dict[int, str] | None = None):
    """Generic styled table: dark header row, alternating row shading.
    header_fill_overrides: {column_index: hex_color} to override the
    default header fill for specific columns (e.g. Ash color for a
    'Seq'/'Task Name' column)."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)
        run.font.name = FONT_NAME
        fill = (header_fill_overrides or {}).get(i, HEADER_FILL)
        _set_cell_shading(hdr_cells[i], fill)

    for r_idx, row_vals in enumerate(rows):
        row_cells = table.add_row().cells
        for c_idx, val in enumerate(row_vals):
            row_cells[c_idx].text = ""
            run = row_cells[c_idx].paragraphs[0].add_run(str(val) if val not in (None, "") else "-")
            run.font.size = Pt(9)
            run.font.name = FONT_NAME
            if r_idx % 2 == 1:
                _set_cell_shading(row_cells[c_idx], ALT_ROW_FILL)

    if col_widths_in:
        for row in table.rows:
            for i, w in enumerate(col_widths_in):
                row.cells[i].width = Inches(w)

    return table


def add_excel_table(doc: Document, headers: list[str], rows: list[list[str]],
                    col_widths_in: list[float] | None = None,
                    group_headers: list[list] | None = None):
    """Render an Excel-sourced rectangular table as an editable Word table.

    Unlike add_table(), blank source cells stay blank rather than becoming
    a dash. When ``group_headers`` is given ([[label, span], ...]) the table
    gets a two-row header: merged group labels on top, column labels below.
    A ["", 1] group means that column's label spans both header rows.
    """
    two_level = bool(group_headers) and sum(g[1] for g in group_headers) == len(headers)
    n_hdr = 2 if two_level else 1
    table = doc.add_table(rows=n_hdr, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False if col_widths_in else True

    def _hdr(cell, text):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text or ""))
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)
        run.font.name = FONT_NAME
        _set_cell_shading(cell, HEADER_FILL)

    if two_level:
        top, sub = table.rows[0].cells, table.rows[1].cells
        c = 0
        for label, span in group_headers:
            if span == 1 and not label:
                # label spans both header rows
                _hdr(top[c], headers[c])
                _set_cell_shading(sub[c], HEADER_FILL)
                top[c].merge(sub[c])
            else:
                for k in range(c, c + span):
                    _hdr(sub[k], headers[k])
                _hdr(top[c], label)
                if span > 1:
                    top[c].merge(top[c + span - 1])
                    top[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c += span
    else:
        for i, h in enumerate(headers):
            _hdr(table.rows[0].cells[i], h)

    for r_idx, row_vals in enumerate(rows):
        cells = table.add_row().cells
        for c_idx in range(len(headers)):
            value = row_vals[c_idx] if c_idx < len(row_vals) else ""
            cells[c_idx].text = ""
            run = cells[c_idx].paragraphs[0].add_run(str(value or ""))
            run.font.size = Pt(9)
            run.font.name = FONT_NAME
            if str(value or "").strip() == "\u2713":
                cells[c_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx % 2 == 1:
                _set_cell_shading(cells[c_idx], ALT_ROW_FILL)

    if col_widths_in:
        for row in table.rows:
            for i, w in enumerate(col_widths_in):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)
    return table


def add_field(paragraph, field_code: str):
    """Insert a Word field (e.g. 'PAGE', 'NUMPAGES') into a paragraph.
    Renders correctly once Word computes it - shows 0/blank in raw XML
    until the user opens the doc (same caveat as add_toc)."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def add_label_value_table(doc: Document, pairs: list[tuple[str, str]]):
    """Borderless 2-column label:value table, used for the Document
    Control Sheet contact block."""
    table = doc.add_table(rows=0, cols=2)
    for label, value in pairs:
        row = table.add_row().cells
        row[0].width = Inches(1.8)
        row[1].width = Inches(4.5)
        run = row[0].paragraphs[0].add_run(f"{label}:")
        run.bold = True
        run.font.name = FONT_NAME
        run2 = row[1].paragraphs[0].add_run(str(value) if value else "")
        run2.font.name = FONT_NAME
    # strip borders for a clean label:value look
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)
    return table


def add_signoff_table(doc: Document, title: str):
    """'Prepared By' / 'Checked By' / 'Approved By' style table:
    header row (Name, Designation, Date, Signature) + one blank row."""
    doc.add_paragraph(title).runs[0].bold = True
    headers = ["Name", "Designation", "Date", "Signature"]
    return add_table(doc, headers, [["", "", "", ""]], col_widths_in=[1.7, 1.7, 1.3, 1.7])


def merge_cells_horizontal(table, row_idx: int, start_col: int, end_col: int):
    """Merge cells in one row from start_col to end_col (inclusive)."""
    row = table.rows[row_idx]
    merged = row.cells[start_col]
    for c in range(start_col + 1, end_col + 1):
        merged = merged.merge(row.cells[c])
    return merged


def add_attribute_values_table(doc: Document, pairs: list[tuple[str, list[str]]]):
    """Two-column 'Attribute' / 'Values' table. When an attribute has
    multiple values, it gets one row per value with the Attribute cell
    vertically merged across them. Attributes with zero captured values
    are skipped entirely by the caller before this is invoked."""
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(("Attribute", "Values")):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)
        run.font.name = FONT_NAME
        _set_cell_shading(hdr_cells[i], HEADER_FILL)

    row_idx = 1
    for attr_name, values in pairs:
        vals = values or ["-"]
        start_row = row_idx
        for v in vals:
            cells = table.add_row().cells
            cells[0].paragraphs[0].text = ""
            # Write the attribute name ONLY in the first row of its group.
            # python-docx's cell.merge() appends the text of every merged
            # cell, so writing it on each row made the name repeat N times.
            if row_idx == start_row:
                run0 = cells[0].paragraphs[0].add_run(attr_name)
                run0.font.size = Pt(9)
                run0.font.name = FONT_NAME
            cells[1].paragraphs[0].text = ""
            run1 = cells[1].paragraphs[0].add_run(str(v))
            run1.font.size = Pt(9)
            run1.font.name = FONT_NAME
            row_idx += 1
        end_row = row_idx - 1
        if end_row > start_row:
            table.cell(start_row, 0).merge(table.cell(end_row, 0))

    return table


def add_toc(doc: Document):
    """Insert a Word TOC field. User must right-click -> Update Field after opening,
    since python-docx/Word don't compute page numbers at generation time."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)
    return paragraph
