"""PNG rendering of the swimlane flowcharts + insertion into the Word guide."""
from __future__ import annotations

import os
import re
from typing import Iterable

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon, Ellipse, Rectangle
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models.schemas import Task
from render import task_flow as tf
from render.task_flow_layout import (Layout, build_layouts, IMG_W, TITLE_H, LANE_HEAD_H, wrap_text,
                                     group_tasks)

DPI = 200
LANE_FILL, LANE_ALT_FILL = "#F4F6F8", "#FFFFFF"
LANE_HEADER_FILL, LANE_EDGE = "#E4E7EB", "#9AA3AD"
TASK_FILL, TASK_EDGE = "#ECEFF2", "#8A939D"
TERM_FILL = "#E4F0F8"
LINE = "#B8962E"
DEC_FILL = "#F6F0DD"
TITLE_FILL = "#FFFFFF"


def draw_layout_png(lay: Layout, out_path: str) -> str:
    fig = plt.figure(figsize=(lay.width, lay.height), dpi=DPI)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, lay.width)
    ax.set_ylim(lay.height, 0)
    ax.axis("off")
    ax.add_patch(Rectangle((0, 0), lay.width, TITLE_H, facecolor=TITLE_FILL, edgecolor=LANE_EDGE,
                           linewidth=0.8, zorder=1))
    ax.text(0.08, TITLE_H / 2, lay.title, ha="left", va="center", fontsize=8.5, fontweight="bold",
            color="#111111", zorder=2)
    for i, lane in enumerate(lay.lanes):
        x = i * lay.lane_w
        ax.add_patch(Rectangle((x, TITLE_H), lay.lane_w, lay.height - TITLE_H,
                               facecolor=LANE_FILL if i % 2 == 0 else LANE_ALT_FILL,
                               edgecolor=LANE_EDGE, linewidth=0.8, zorder=0))
        ax.add_patch(Rectangle((x, TITLE_H), lay.lane_w, LANE_HEAD_H, facecolor=LANE_HEADER_FILL,
                               edgecolor=LANE_EDGE, linewidth=0.8, zorder=1))
        txt, _ = wrap_text(tf.lane_title(lane), lay.lane_w - 0.1, 7.0)
        ax.text(x + lay.lane_w / 2, TITLE_H + LANE_HEAD_H / 2, txt, ha="center", va="center",
                fontsize=7.0, fontweight="bold", color="#222222", zorder=2)

    for nd in lay.nodes:
        if nd.kind == "task":
            ax.add_patch(FancyBboxPatch((nd.left, nd.top), nd.w, nd.h, boxstyle="round,pad=0,rounding_size=0.04",
                                        facecolor=TASK_FILL, edgecolor=TASK_EDGE, linewidth=0.9, zorder=3))
            ax.text(nd.cx, nd.cy, nd.label, ha="center", va="center", fontsize=nd.font, color="#1B1B1B",
                    zorder=5, linespacing=1.05)
        elif nd.kind == "decision":
            ax.add_patch(Polygon([(nd.cx, nd.top), (nd.right, nd.cy), (nd.cx, nd.bottom), (nd.left, nd.cy)],
                                 closed=True, facecolor=DEC_FILL, edgecolor=LINE, linewidth=1.0, zorder=3))
            ax.text(nd.cx, nd.cy, nd.label, ha="center", va="center", fontsize=nd.font, color="#1B1B1B",
                    zorder=5, linespacing=1.0)
        else:
            ax.add_patch(Ellipse((nd.cx, nd.cy), nd.w, nd.h, facecolor=TERM_FILL, edgecolor=TASK_EDGE,
                                 linewidth=0.9, zorder=3))
            ax.text(nd.cx, nd.cy, nd.label, ha="center", va="center", fontsize=nd.font, zorder=5)

    arrow = dict(arrowstyle="-|>", color=LINE, lw=1.0, mutation_scale=7, shrinkA=0, shrinkB=0)
    for e in lay.edges:
        xs, ys = [p[0] for p in e.points], [p[1] for p in e.points]
        ax.plot(xs, ys, color=LINE, linewidth=1.0, zorder=2, solid_capstyle="round")
        if e.heads in ("end", "both"):
            ax.annotate("", xy=e.points[-1], xytext=e.points[-2], arrowprops=arrow, zorder=2)
        if e.heads == "both":
            ax.annotate("", xy=e.points[0], xytext=e.points[1], arrowprops=arrow, zorder=2)
        if e.label:
            ax.text(e.label_xy[0], e.label_xy[1], e.label, fontsize=6, color="#7A6420", ha="left",
                    va="top" if e.label_below else "bottom", zorder=6)
    if lay.total > 1:
        ax.text(lay.width - 0.05, lay.height - 0.06, f"Part {lay.panel_no} of {lay.total}", ha="right",
                va="bottom", fontsize=6, color="#555555", zorder=6)
    fig.savefig(out_path, dpi=DPI, facecolor="white")
    plt.close(fig)
    return out_path


def render_flow_images_ex(tasks: Iterable[Task], out_dir: str, name: str, product: str = ""):
    """-> (paths, scales, warnings). scale < 1 means the picture is inserted slightly smaller
    in Word so the whole diagram fits on one page."""
    layouts, warnings = build_layouts(tasks, name, product)
    if not layouts:
        return [], [], warnings
    os.makedirs(out_dir, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_") or "flow"
    paths = []
    for lay in layouts:
        path = os.path.join(out_dir, f"flow_{safe}_{lay.panel_no}.png")
        draw_layout_png(lay, path)
        paths.append(path)
    return paths, [lay.scale for lay in layouts], warnings


def render_flow_images(tasks: Iterable[Task], out_dir: str, name: str, product: str = ""):
    """Create one PNG per diagram part for the given order type. Returns (paths, warnings)."""
    paths, _scales, warnings = render_flow_images_ex(tasks, out_dir, name, product)
    return paths, warnings


def add_task_flowchart_image(doc: Document, tasks: Iterable[Task], order_type: str, out_dir: str,
                             product: str = "") -> dict:
    """Render the swimlane diagram(s) for one order type and insert them."""
    tasks = list(tasks)
    paths, scales, warnings = render_flow_images_ex(tasks, out_dir, order_type, product)
    if not paths:
        p = doc.add_paragraph()
        r = p.add_run("[No eligible tasks for flowchart: External/System tasks were excluded, "
                      "or no swimlane could be assigned.]")
        r.italic = True
        return {"eligible_tasks": 0, "panels": 0, "lanes": [], "images": [], "warnings": warnings}
    flow_tasks = tf.prepare_flow_tasks(tasks)
    from render.task_flow_layout import lane_order
    lanes = lane_order(group_tasks(flow_tasks))
    for i, path in enumerate(paths, start=1):
        if i > 1:
            doc.add_page_break()
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(path, width=Inches(IMG_W * scales[i - 1]))
    return {"eligible_tasks": len(flow_tasks), "panels": len(paths), "lanes": lanes,
            "images": paths, "warnings": warnings}
