"""Editable PowerPoint export of the swimlane flowcharts.

One slide per diagram panel, drawn from the same Layout as the PNG images.
Everything is a native PowerPoint object (shapes, text, connectors), so a
user can retitle a task, move a box, re-route or delete an arrow, and so on.
Arrows are real connectors glued to the boxes wherever the route allows
(straight and elbow connectors move with the shapes); complex loop-around
routes are drawn as editable freeform lines instead.
"""
from __future__ import annotations

from typing import Iterable

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from render import task_flow as tf
from render.task_flow_layout import (Layout, build_layouts, TITLE_H, LANE_HEAD_H)

SLIDE_W, SLIDE_H = 8.27, 11.69          # A4 portrait, matches the Word page
OFF_X, OFF_Y = (SLIDE_W - 6.2) / 2, 0.4

C_LANE, C_LANE_ALT = RGBColor(0xF4, 0xF6, 0xF8), RGBColor(0xFF, 0xFF, 0xFF)
C_LANE_HEAD, C_LANE_EDGE = RGBColor(0xE4, 0xE7, 0xEB), RGBColor(0x9A, 0xA3, 0xAD)
C_TASK, C_TASK_EDGE = RGBColor(0xEC, 0xEF, 0xF2), RGBColor(0x8A, 0x93, 0x9D)
C_TERM = RGBColor(0xE4, 0xF0, 0xF8)
C_DEC, C_LINE = RGBColor(0xF6, 0xF0, 0xDD), RGBColor(0xB8, 0x96, 0x2E)
C_TEXT = RGBColor(0x1B, 0x1B, 0x1B)
TOL = 0.012


def _emu(v: float) -> int:
    return int(round(v * 914400))


def _X(x):
    return _emu(OFF_X + x)


def _Y(y):
    return _emu(OFF_Y + y)


def _style(shape, fill, line, width=0.9):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    shape.shadow.inherit = False


def _text(shape, text, size, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, color=C_TEXT):
    tf_ = shape.text_frame
    tf_.word_wrap = True
    tf_.vertical_anchor = anchor
    tf_.margin_left = tf_.margin_right = Inches(0.03)
    tf_.margin_top = tf_.margin_bottom = Inches(0.01)
    lines = str(text).split("\n")
    # PowerPoint wraps by itself, so join the pre-wrapped lines back into one paragraph
    p = tf_.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = " ".join(lines)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Calibri"
    r.font.color.rgb = color


def _arrows(conn, heads):
    ln = conn.line._get_or_add_ln()
    for tag in ("a:headEnd", "a:tailEnd"):
        for old in ln.findall(qn(tag)):
            ln.remove(old)
    if heads == "both":
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if heads in ("end", "both"):
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")


def _line_style(conn):
    conn.line.color.rgb = C_LINE
    conn.line.width = Pt(1.0)
    try:
        conn.shadow.inherit = False
    except Exception:
        pass


def _site(node, pt):
    """Connection-site index (0 top, 1 left, 2 bottom, 3 right) if pt is on it."""
    sites = {0: (node.cx, node.top), 1: (node.left, node.cy), 2: (node.cx, node.bottom), 3: (node.right, node.cy)}
    for idx, (sx, sy) in sites.items():
        if abs(sx - pt[0]) < TOL and abs(sy - pt[1]) < TOL:
            return idx
    return None


def _glue(conn, which, shape, idx):
    c = conn._element.nvCxnSpPr.cNvCxnSpPr
    tag = qn("a:stCxn") if which == "start" else qn("a:endCxn")
    for old in c.findall(tag):
        c.remove(old)
    el = etree.SubElement(c, tag, id=str(shape.shape_id), idx=str(idx))
    if which == "start":      # schema order: stCxn before endCxn
        c.remove(el)
        c.insert(0, el)


def _simplify(points):
    pts = [points[0]]
    for p in points[1:]:
        if abs(p[0] - pts[-1][0]) > 1e-6 or abs(p[1] - pts[-1][1]) > 1e-6:
            pts.append(p)
    out = [pts[0]]
    for i in range(1, len(pts) - 1):
        a, b, c = out[-1], pts[i], pts[i + 1]
        if (abs(a[0] - b[0]) < 1e-6 and abs(b[0] - c[0]) < 1e-6) or (abs(a[1] - b[1]) < 1e-6 and abs(b[1] - c[1]) < 1e-6):
            continue
        out.append(b)
    out.append(pts[-1])
    return out


def _set_xfrm(conn, x, y, cx, cy, rot=0, flip_h=False, flip_v=False):
    xfrm = conn._element.spPr.find(qn("a:xfrm"))
    for a in ("rot", "flipH", "flipV"):
        xfrm.attrib.pop(a, None)
    if rot:
        xfrm.set("rot", str(rot))
    if flip_h:
        xfrm.set("flipH", "1")
    if flip_v:
        xfrm.set("flipV", "1")
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    off.set("x", str(x)), off.set("y", str(y))
    ext.set("cx", str(cx)), ext.set("cy", str(cy))


def _set_geom(conn, prst, adj=None):
    geom = conn._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", prst)
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for g in list(av):
        av.remove(g)
    if adj is not None:
        etree.SubElement(av, qn("a:gd"), name="adj1", fmla=f"val {int(round(adj * 100000))}")


def _rotated_elbow(slide, p0, p1, kind, adj=0.5):
    """V-first elbow going DOWN from p0 to p1 (rot 90 deg). kind: 'L' or 'Z'."""
    (x0, y0), (x1, y1) = p0, p1
    dx, dy = abs(x1 - x0), abs(y1 - y0)
    conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, _X(x0), _Y(y0), _X(x1), _Y(y1))
    cxm, cym = (x0 + x1) / 2, (y0 + y1) / 2
    ext_u, ext_v = _emu(dy), _emu(dx)           # local u = vertical extent after rotation
    _set_xfrm(conn, _X(cxm) - ext_u // 2, _Y(cym) - ext_v // 2, ext_u, ext_v,
              rot=5400000, flip_v=(x1 > x0))
    _set_geom(conn, "bentConnector2" if kind == "L" else "bentConnector3", None if kind == "L" else adj)
    return conn


def _add_edge(slide, e, nodes, shape_of):
    pts = _simplify(e.points)
    first, last = pts[0], pts[-1]
    conn = None
    n = len(pts)
    vert = lambda a, b: abs(a[0] - b[0]) < 1e-6
    horiz = lambda a, b: abs(a[1] - b[1]) < 1e-6
    try:
        if n == 2:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _X(first[0]), _Y(first[1]),
                                              _X(last[0]), _Y(last[1]))
        elif n == 3 and vert(pts[0], pts[1]) and horiz(pts[1], pts[2]) and pts[2][1] > pts[0][1]:
            conn = _rotated_elbow(slide, pts[0], pts[2], "L")
        elif n == 3 and horiz(pts[0], pts[1]) and vert(pts[1], pts[2]):
            conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, _X(first[0]), _Y(first[1]),
                                              _X(last[0]), _Y(last[1]))
            _set_geom(conn, "bentConnector2")
        elif n == 4 and vert(pts[0], pts[1]) and horiz(pts[1], pts[2]) and vert(pts[2], pts[3]) \
                and pts[3][1] > pts[0][1] and abs(pts[3][0] - pts[0][0]) > 0.02:
            adj = (pts[1][1] - pts[0][1]) / (pts[3][1] - pts[0][1])
            if 0.02 < adj < 0.98:
                conn = _rotated_elbow(slide, pts[0], pts[3], "Z", adj)
        elif n == 4 and horiz(pts[0], pts[1]) and vert(pts[1], pts[2]) and horiz(pts[2], pts[3]) \
                and abs(pts[3][0] - pts[0][0]) > 0.02 and abs(pts[3][1] - pts[0][1]) > 0.02:
            adj = (pts[1][0] - pts[0][0]) / (pts[3][0] - pts[0][0])
            if 0.02 < adj < 0.98:
                conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, _X(first[0]), _Y(first[1]),
                                                  _X(last[0]), _Y(last[1]))
                _set_geom(conn, "bentConnector3", adj)
    except Exception:
        conn = None
    glue = conn is not None
    if conn is None:                              # editable freeform polyline
        fb = slide.shapes.build_freeform(_X(pts[0][0]), _Y(pts[0][1]))
        fb.add_line_segments([(_X(x), _Y(y)) for x, y in pts[1:]], close=False)
        conn = fb.convert_to_shape()
        conn.fill.background()
    _line_style(conn)
    _arrows(conn, e.heads)
    if glue:
        for which, pt in (("start", first), ("end", last)):
            for nd in nodes:
                idx = _site(nd, pt)
                if idx is not None and id(nd) in shape_of:
                    _glue(conn, which, shape_of[id(nd)], idx)
                    break
    if e.label:
        tb = slide.shapes.add_textbox(_X(e.label_xy[0]) - _emu(0.02), _Y(e.label_xy[1]) - _emu(0.13),
                                      _emu(0.3), _emu(0.14))
        tb.text_frame.margin_left = tb.text_frame.margin_right = 0
        tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = e.label
        r.font.size = Pt(6)
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(0x7A, 0x64, 0x20)
        tb.name = f"Label {e.label}"
    return conn


def add_layout_slide(prs: Presentation, lay: Layout):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # title bar
    tb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _X(0), _Y(0), _emu(lay.width), _emu(TITLE_H))
    _style(tb, RGBColor(0xFF, 0xFF, 0xFF), C_LANE_EDGE)
    _text(tb, lay.title, 9, bold=True, align=PP_ALIGN.LEFT, color=RGBColor(0x11, 0x11, 0x11))
    tb.name = "Title"
    # swimlanes
    for i, lane in enumerate(lay.lanes):
        x = i * lay.lane_w
        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _X(x), _Y(TITLE_H), _emu(lay.lane_w),
                                      _emu(lay.height - TITLE_H))
        _style(body, C_LANE if i % 2 == 0 else C_LANE_ALT, C_LANE_EDGE, 0.75)
        body.name = f"Lane body {tf.lane_title(lane)}"
        head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _X(x), _Y(TITLE_H), _emu(lay.lane_w),
                                      _emu(LANE_HEAD_H))
        _style(head, C_LANE_HEAD, C_LANE_EDGE, 0.75)
        _text(head, tf.lane_title(lane), 7.5, bold=True)
        head.name = f"Lane {tf.lane_title(lane)}"
    # nodes
    shape_of = {}
    for nd in lay.nodes:
        kind = {"task": MSO_SHAPE.ROUNDED_RECTANGLE, "decision": MSO_SHAPE.FLOWCHART_DECISION,
                "oval": MSO_SHAPE.FLOWCHART_TERMINATOR}[nd.kind]
        shp = slide.shapes.add_shape(kind, _X(nd.left), _Y(nd.top), _emu(nd.w), _emu(nd.h))
        if nd.kind == "decision":
            _style(shp, C_DEC, C_LINE, 1.0)
        elif nd.kind == "oval":
            _style(shp, C_TERM, C_TASK_EDGE, 0.9)
        else:
            _style(shp, C_TASK, C_TASK_EDGE, 0.9)
            try:
                shp.adjustments[0] = 0.1
            except Exception:
                pass
        _text(shp, nd.label, nd.font + (0.6 if nd.kind != "decision" else 0.2))
        shp.name = f"{nd.kind.title()}: {nd.label.splitlines()[0] if nd.label else ''}"
        shape_of[id(nd)] = shp
    for e in lay.edges:
        _add_edge(slide, e, lay.nodes, shape_of)
    return slide


def build_flowcharts_pptx(sections: Iterable, out_path: str, product: str = "") -> dict:
    """sections: iterable of objects with .order_type and .tasks. Returns a summary dict."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    warnings, slides = [], 0
    for sec in sections:
        layouts, warns = build_layouts(sec.tasks, sec.order_type, product)
        warnings += [f"[{sec.order_type}] {w}" for w in warns]
        for lay in layouts:
            add_layout_slide(prs, lay)
            slides += 1
    if slides == 0:
        raise ValueError("No eligible tasks for any flowchart - nothing to export.")
    prs.save(out_path)
    return {"slides": slides, "path": out_path, "warnings": warnings}
