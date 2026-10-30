"""Write the test plan as an Excel workbook ready for UAT execution."""
from __future__ import annotations

import datetime
import io

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from tools.testx.model import BLANK, BLANK_LABEL
from tools.testx.plan import Plan, safe_sheet_name

NAVY, BLUE, TEAL, PALE, GREEN = "0B2D55", "0879D9", "00A6B2", "EAF1F8", "19C995"
FONT = "Calibri"
STATUSES = ["Not Run", "Pass", "Fail", "Blocked"]
EXEC_COLS = ["Rules", "Expected Result", "Actual Result", "Status", "Tested By", "Test Date", "Remarks"]

thin = Side(style="thin", color="CFDAE8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
EXEC_FILL = PatternFill("solid", fgColor=BLUE)
PALE_FILL = PatternFill("solid", fgColor=PALE)
NEG_FILL = PatternFill("solid", fgColor="FFF4DB")
BAD_FILL = PatternFill("solid", fgColor="FFD9D4")


def _f(bold=False, color="000000", size=11, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _shown(v):
    return BLANK_LABEL if v is BLANK else v


def build_workbook(plan: Plan, source_name: str = "", prepared_by: str = "") -> bytes:
    wb = Workbook()
    used: set = set()
    summary = wb.active
    summary.title = safe_sheet_name("Summary", used)
    sheet_names = []
    for g in plan.groups:
        sheet_names.append(safe_sheet_name(g.name, used))
    rules_name = safe_sheet_name("Rules", used)

    # ---------------------------------------------------------------- group sheets
    for g, sname in zip(plan.groups, sheet_names):
        ws = wb.create_sheet(sname)
        attr_keys = [f.key for f in g.factors]
        headers = ["Test ID", "Type"] + attr_keys + EXEC_COLS
        n_cols = len(headers)
        ws["A1"] = f"{plan.product} · {g.name}"
        ws["A1"].font = _f(True, NAVY, 15)
        pos = sum(1 for c in g.cases if c.kind == "Positive")
        neg = len(g.cases) - pos
        technique = "pairwise (every pair of values covered)" if plan.settings.technique == "pairwise" else "all valid combinations"
        ws["A2"] = f"{pos} positive and {neg} negative test(s) · technique: {technique}"
        ws["A2"].font = _f(False, "63758B", 10)
        if g.free_entry:
            ws["A3"] = "Not combined (no list of values - tester enters a valid value): " + ", ".join(
                f"{a}{' *' if m else ''}" for a, _n, m in g.free_entry) + "   (* mandatory)"
            ws["A3"].font = _f(False, "63758B", 10, True)
        hr = 5
        for j, h in enumerate(headers, start=1):
            c = ws.cell(hr, j, h)
            c.font = _f(True, "FFFFFF", 10)
            c.fill = EXEC_FILL if h in EXEC_COLS and h != "Rules" else HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
        ws.row_dimensions[hr].height = 32
        widths = {h: max(len(str(h)) + 2, 10) for h in headers}
        status_col = headers.index("Status") + 1
        for r, case in enumerate(g.cases, start=hr + 1):
            vals = [case.id, case.kind] + [_shown(case.values.get(k)) for k in attr_keys] + [
                ", ".join(case.rules), case.expected, "", "Not Run", "", "", ""]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(r, j, v)
                c.font = _f(size=10)
                c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if case.kind == "Negative":
                    c.fill = NEG_FILL
                name = headers[j - 1]
                if name in attr_keys and v == BLANK_LABEL:
                    c.font = _f(size=10, color="8A9BB0", italic=True)
                if isinstance(v, str) and name not in ("Expected Result", "Actual Result", "Remarks"):
                    widths[name] = min(max(widths[name], len(v) + 2), 30)
            if case.highlight and case.highlight in attr_keys:
                bad = ws.cell(r, 3 + attr_keys.index(case.highlight))
                bad.fill = BAD_FILL
                bad.font = _f(True, "B42318", 10)
            ws.cell(r, 1).font = _f(True, NAVY, 10)
        last = hr + len(g.cases)
        for j, h in enumerate(headers, start=1):
            w = {"Expected Result": 52, "Actual Result": 28, "Remarks": 28, "Tested By": 16, "Test Date": 13,
                 "Status": 12, "Rules": 14}.get(h, widths[h])
            ws.column_dimensions[get_column_letter(j)].width = w
        if g.cases:
            dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(status_col)}{hr + 1}:{get_column_letter(status_col)}{last}")
            rng = f"{get_column_letter(status_col)}{hr + 1}:{get_column_letter(status_col)}{last}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", bgColor="C6F0DF")))
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", bgColor="FFD0CA")))
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", bgColor="FFE9A8")))
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(n_cols)}{last}"
        ws.freeze_panes = ws.cell(hr + 1, 3)
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = BLUE

    # --------------------------------------------------------------------- rules sheet
    wr = wb.create_sheet(rules_name)
    wr["A1"] = "Business rules used to build the plan"
    wr["A1"].font = _f(True, NAVY, 15)
    wr["A2"] = "Rules were read from the workbook (sheet 3) or added in TestX. A rule marked 'Not used' was switched off."
    wr["A2"].font = _f(False, "63758B", 10)
    heads = ["Rule", "Origin", "Rule text", "Used", "Status", "Notes", "Applied on sheet(s)"]
    for j, h in enumerate(heads, start=1):
        c = wr.cell(4, j, h)
        c.font, c.fill, c.border = _f(True, "FFFFFF", 10), HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    applied = {}
    for g, sname in zip(plan.groups, sheet_names):
        for rid in g.rules_applied:
            applied.setdefault(rid, []).append(sname)
    for r, rule in enumerate(plan.rules, start=5):
        vals = [rule.id, {"table": "Table in workbook", "sentence": "Sentence in workbook", "manual": "Added in TestX"}.get(rule.source, rule.source),
                rule.text, "Yes" if rule.enabled else "Not used", rule.status,
                "; ".join(rule.issues + rule.notes), ", ".join(applied.get(rule.id, []))]
        for j, v in enumerate(vals, start=1):
            c = wr.cell(r, j, v)
            c.font, c.border = _f(size=10), BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if not rule.enabled:
            for j in range(1, len(heads) + 1):
                wr.cell(r, j).font = _f(size=10, color="8A9BB0", italic=True)
    for j, w in enumerate([8, 20, 80, 10, 14, 50, 28], start=1):
        wr.column_dimensions[get_column_letter(j)].width = w
    wr.freeze_panes = "A5"
    wr.sheet_properties.tabColor = TEAL

    # ------------------------------------------------------------------------ summary
    s = summary
    s["A1"] = "TestX · UAT Test Plan"
    s["A1"].font = _f(True, NAVY, 20)
    s["A2"] = plan.product
    s["A2"].font = _f(True, BLUE, 13)
    meta = [
        ("Generated", datetime.datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Source workbook", source_name or "-"),
        ("Prepared by", prepared_by or "-"),
        ("Technique", "Pairwise - every pair of attribute values is covered at least once" if plan.settings.technique == "pairwise"
         else "All valid combinations"),
        ("Scope", "Each component on its own sheet" if plan.settings.scope == "component" else "All selected components together"),
        ("Negative tests", "Included - one per rule" if plan.settings.negatives else "Not included"),
        ("Order type varied", "Yes" if plan.settings.vary_order_type else "No"),
    ]
    for i, (k, v) in enumerate(meta, start=4):
        s.cell(i, 1, k).font = _f(True, NAVY, 10)
        s.cell(i, 2, v).font = _f(size=10)
        s.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
    r0 = 4 + len(meta) + 1
    heads = ["Sheet", "Positive tests", "Negative tests", "Total", "All combinations", "Value pairs", "Pairs that cannot occur"]
    for j, h in enumerate(heads, start=1):
        c = s.cell(r0, j, h)
        c.font, c.fill, c.border = _f(True, "FFFFFF", 10), HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    s.row_dimensions[r0].height = 30
    for i, (g, sname) in enumerate(zip(plan.groups, sheet_names), start=r0 + 1):
        pos = sum(1 for c in g.cases if c.kind == "Positive")
        row = [sname, pos, len(g.cases) - pos, len(g.cases), g.full_size,
               g.pairs_total if plan.settings.technique == "pairwise" else "-",
               g.pairs_unreachable if plan.settings.technique == "pairwise" else "-"]
        for j, v in enumerate(row, start=1):
            c = s.cell(i, j, v)
            c.font, c.border = _f(size=10), BORDER
            if j == 1:
                c.hyperlink = f"#'{sname}'!A1"
                c.font = Font(name=FONT, size=10, color=BLUE, underline="single")
            else:
                c.alignment = Alignment(horizontal="right")
                if isinstance(v, int):
                    c.number_format = "#,##0"
    tot = r0 + 1 + len(plan.groups)
    s.cell(tot, 1, "Total").font = _f(True, NAVY, 10)
    for j, col in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        c = s.cell(tot, j, f"=SUM({col}{r0 + 1}:{col}{tot - 1})" if plan.groups else 0)
        c.font, c.number_format = _f(True, NAVY, 10), "#,##0"
        c.alignment = Alignment(horizontal="right")
    for j in range(1, 8):
        s.cell(tot, j).fill = PALE_FILL
        s.cell(tot, j).border = BORDER
    link = s.cell(tot + 2, 1, "Business rules used")
    link.hyperlink = f"#'{rules_name}'!A1"
    link.font = Font(name=FONT, size=10, color=BLUE, underline="single")
    notes = list(plan.warnings)
    for g in plan.groups:
        notes += [f"{g.name}: {n}" for n in g.notes]
    if notes:
        s.cell(tot + 4, 1, "Notes").font = _f(True, NAVY, 11)
        for i, n in enumerate(notes, start=tot + 5):
            c = s.cell(i, 1, n)
            c.font = _f(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            s.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)
            s.row_dimensions[i].height = max(15, 15 * (len(n) // 95 + 1))
    for j, w in enumerate([30, 16, 16, 12, 18, 14, 20], start=1):
        s.column_dimensions[get_column_letter(j)].width = w
    s.column_dimensions["B"].width = 18
    s.sheet_view.showGridLines = False
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.sheet_properties.tabColor = NAVY

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
