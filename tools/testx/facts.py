"""Extra facts TestX needs for scenario-style test cases.

The components, order types, tasks and special requirements are still read by the existing GuideX
extractors (see source.py) - nothing about them changes. This module only adds small, tolerant
readers for sheets and columns that the document generator does not need:

    2. Product Hierarchy   Entity Mapping / Parent / Mandatory? / Comments of every product row
    3. Special Requirements  Requirement ID, rule type, applicable system
    4. Pricing             price plans with override ranges
    5.1 Order Scenarios    order sub types / conditions per order type
    5.2 Attribute vs Order attributes changed by a modification
    10.0 Fault Handling    work groups and escalation levels

Every reader returns an empty list instead of failing, so a workbook that lacks a sheet simply gets
fewer scenarios.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl


def _s(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _key(v) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _s(v).lower()).strip()


def _find_sheet(wb, *needles: str):
    """First sheet whose lower-case name contains every needle."""
    for name in wb.sheetnames:
        low = name.lower()
        if all(n in low for n in needles):
            return wb[name]
    return None


def _header_row(ws, wanted: list[str], scan: int = 25):
    """(row, {wanted_key: column}) for the first row that contains most of the wanted headers."""
    best = (None, {})
    for r in range(1, min(scan, ws.max_row) + 1):
        cols = {}
        for c in range(1, min(ws.max_column, 60) + 1):
            k = _key(ws.cell(r, c).value)
            if not k:
                continue
            for w in wanted:
                if w not in cols and (k == w or k.startswith(w)):
                    cols[w] = c
        if len(cols) > len(best[1]):
            best = (r, cols)
    return best


# ---------------------------------------------------------------------------- data
@dataclass
class HierExtra:
    mapping: str = ""          # "CRM, OSS SA, BSS"
    parent: str = ""
    mandatory: str = ""
    comments: str = ""
    charge: str = ""           # MRC / NRC


@dataclass
class PriceLine:
    product: str
    plan: str
    related: str
    charge: str
    override: bool
    low: str
    high: str
    period: str
    prorated: str
    parametric: str


@dataclass
class FaultRow:
    product: str
    work_group: str
    ola: str
    nature: str
    cause: str
    l1: str
    l2: str
    l3: str


@dataclass
class ModifyRow:
    kind: str                  # ATTRIBUTE | FEATURE
    product: str
    name: str
    back_update: bool
    action: str
    order_type: str


@dataclass
class OrderRow:
    sub_type: str
    condition: str
    action: str
    order_type: str
    status_remarks: str


@dataclass
class Requirement:
    rid: str
    rtype: str
    systems: str
    text: str


@dataclass
class ValueNote:
    product: str
    attribute: str
    value: str
    comment: str


@dataclass
class Extra:
    hier: dict = field(default_factory=dict)          # source_row -> HierExtra
    value_notes: list = field(default_factory=list)   # comments written on individual Value rows
    pricing: list = field(default_factory=list)
    faults: list = field(default_factory=list)
    modify_rows: list = field(default_factory=list)
    order_rows: list = field(default_factory=list)
    requirements: list = field(default_factory=list)
    task_sheet: str = ""


# ------------------------------------------------------------------------- readers
def find_task_sheet(names: list[str]) -> str:
    """The sheet that holds the stacked order-type task lists (6.2 or 6.3 depending on the template)."""
    for pref in ("so task list", "task list"):
        for n in names:
            if pref in n.lower():
                return n
    return ""


def read_hierarchy_extra(wb, sheet_name: str, rows_needed: set[int]) -> dict:
    ws = wb[sheet_name]
    hr, cols = _header_row(ws, ["entity mapping", "parent reference", "mandatory", "comments", "nrc mrc usage"])
    out = {}
    if not hr:
        return out
    g = lambda row, k: _s(ws.cell(row, cols[k]).value) if k in cols else ""
    for r in rows_needed:
        out[r] = HierExtra(mapping=g(r, "entity mapping"), parent=g(r, "parent reference"), mandatory=g(r, "mandatory"),
                           comments=str(ws.cell(r, cols["comments"]).value or "").strip() if "comments" in cols else "",
                           charge=g(r, "nrc mrc usage"))
    return out


def read_value_notes(wb, sheet_name: str) -> list:
    """Comments written on a single Value row, e.g. "This LOV should select as default value when 'SASE Service' is set to 'SASE Fabric 361'."."""
    ws = wb[sheet_name]
    hr, cols = _header_row(ws, ["comments"])
    if not hr or "comments" not in cols:
        return []
    out, product, attr = [], "", ""
    for r in range(hr + 1, ws.max_row + 1):
        kind, name = _key(ws.cell(r, 3).value), _s(ws.cell(r, 4).value)
        if kind in ("product", "product component"):
            product, attr = name, ""
        elif kind == "attribute":
            attr = name
        elif kind == "value":
            note = str(ws.cell(r, cols["comments"]).value or "").strip()
            if note:
                out.append(ValueNote(product, attr, name, note))
    return out


def read_requirements(wb) -> list:
    ws = _find_sheet(wb, "special", "requirement")
    if ws is None:
        return []
    hr, cols = _header_row(ws, ["requirement id", "business rule type", "applicable system", "requirement description"])
    if not hr or "requirement description" not in cols:
        return []
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        text = str(ws.cell(r, cols["requirement description"]).value or "").strip()
        if not text:
            continue
        out.append(Requirement(_s(ws.cell(r, cols.get("requirement id", 1)).value) or f"R{len(out) + 1:02d}",
                               _s(ws.cell(r, cols["business rule type"]).value) if "business rule type" in cols else "",
                               _s(ws.cell(r, cols["applicable system"]).value) if "applicable system" in cols else "", text))
    return out


def read_pricing(wb) -> list:
    ws = _find_sheet(wb, "pricing")
    if ws is None:
        return []
    hr, cols = _header_row(ws, ["main product name", "product name", "price plan name", "price related attributes",
                                "charge type", "parametric", "override", "charge period", "prorated"])
    if not hr or "price plan name" not in cols:
        return []
    lo = hi = None                                   # Min / Max sit in a sub-header row under "Price override"
    for r in range(hr, min(hr + 3, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            k = _key(ws.cell(r, c).value)
            if k == "min" and lo is None and c >= cols.get("override", 1):
                lo = c
            if k == "max" and hi is None and c >= cols.get("override", 1):
                hi = c
    data_start = hr + 1
    while data_start <= ws.max_row and not _s(ws.cell(data_start, cols["price plan name"]).value):
        data_start += 1
    out, product = [], ""
    for r in range(data_start, ws.max_row + 1):
        first = _s(ws.cell(r, 1).value)
        if first.lower().startswith(("taxes", "otc")):
            break
        plan = _s(ws.cell(r, cols["price plan name"]).value)
        if not plan:
            continue
        pn = _s(ws.cell(r, cols["product name"]).value) if "product name" in cols else ""
        product = pn or product
        ov = _s(ws.cell(r, cols["override"]).value).lower() if "override" in cols else ""
        out.append(PriceLine(product, plan, _s(ws.cell(r, cols["price related attributes"]).value) if "price related attributes" in cols else "",
                             _s(ws.cell(r, cols["charge type"]).value).upper() if "charge type" in cols else "",
                             ov.startswith("y"), _s(ws.cell(r, lo).value) if lo else "", _s(ws.cell(r, hi).value) if hi else "",
                             _s(ws.cell(r, cols["charge period"]).value) if "charge period" in cols else "",
                             _s(ws.cell(r, cols["prorated"]).value) if "prorated" in cols else "",
                             _s(ws.cell(r, cols["parametric"]).value) if "parametric" in cols else ""))
    return out


def read_faults(wb) -> list:
    ws = _find_sheet(wb, "fault")
    if ws is None:
        return []
    hr, cols = _header_row(ws, ["product name", "default work group", "ola", "nature of fault", "cause of fault",
                                "level 1", "level 2", "level 3"])
    if not hr or "default work group" not in cols:
        return []
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        pn = _s(ws.cell(r, cols.get("product name", 1)).value)
        wg = _s(ws.cell(r, cols["default work group"]).value)
        if not pn or not wg:
            continue
        g = lambda k: _s(ws.cell(r, cols[k]).value) if k in cols else ""
        out.append(FaultRow(pn, wg, g("ola"), g("nature of fault"), g("cause of fault"), g("level 1"), g("level 2"), g("level 3")))
    return out


def read_modify_rows(wb) -> list:
    ws = _find_sheet(wb, "attribute", "order")
    if ws is None:
        return []
    hr, cols = _header_row(ws, ["input type", "product name", "input value", "back update required", "action", "oss order type"])
    if not hr or "input value" not in cols:
        return []
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        first = _s(ws.cell(r, 1).value)
        if first.lower().startswith("combined"):
            break
        kind = _s(ws.cell(r, cols.get("input type", 2)).value).upper()
        name = _s(ws.cell(r, cols["input value"]).value)
        if not kind or not name:
            continue
        out.append(ModifyRow(kind, _s(ws.cell(r, cols.get("product name", 3)).value), name,
                             _s(ws.cell(r, cols.get("back update required", 6)).value).lower().startswith("y"),
                             _s(ws.cell(r, cols.get("action", 8)).value).upper(),
                             _s(ws.cell(r, cols.get("oss order type", 9)).value).upper()))
    return out


def read_order_rows(wb) -> list:
    ws = _find_sheet(wb, "order", "scenario")
    if ws is None:
        return []
    hr, cols = _header_row(ws, ["order sub type", "order condition", "action", "order type list", "value in status remarks"])
    if not hr or "order type list" not in cols:
        return []
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        ot = _s(ws.cell(r, cols["order type list"]).value)
        if not ot:
            continue
        out.append(OrderRow(_s(ws.cell(r, cols.get("order sub type", 4)).value), _s(ws.cell(r, cols.get("order condition", 5)).value),
                            _s(ws.cell(r, cols.get("action", 6)).value), ot.upper(),
                            _s(ws.cell(r, cols["value in status remarks"]).value) if "value in status remarks" in cols else ""))
    return out


def read_extra(workbook_path: str, hierarchy_nodes, hierarchy_sheet_hint: str = "product hierarchy") -> Extra:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ex = Extra()
    ex.task_sheet = find_task_sheet(wb.sheetnames)
    hs = next((n for n in wb.sheetnames if hierarchy_sheet_hint in n.lower()), None)
    if hs and hierarchy_nodes:
        rows = set()
        for n in hierarchy_nodes:
            rows.add(n.source_row)
            rows.update(a.source_row for a in n.attributes)
        ex.hier = read_hierarchy_extra(wb, hs, rows)
        ex.value_notes = read_value_notes(wb, hs)
    ex.requirements = read_requirements(wb)
    ex.pricing = read_pricing(wb)
    ex.faults = read_faults(wb)
    ex.modify_rows = read_modify_rows(wb)
    ex.order_rows = read_order_rows(wb)
    return ex
