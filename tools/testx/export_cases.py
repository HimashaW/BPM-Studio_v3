"""Write the scenario test cases in the layout testers already use:

  Test Case ID | Service Type | Order Type | IT System | Product | Test Scenario |
  Expected result in CRM/OSS/BSS (CRM | OSS | BSS) | Status | Workbook reference
"""
from __future__ import annotations

import datetime
import io

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, BLUE, PALE = "0B2D55", "0879D9", "EAF1F8"
STATUSES = ["Not Run", "Pass", "Fail", "Blocked"]
thin = Side(style="thin", color="B7C6D8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor=NAVY)
SUB = PatternFill("solid", fgColor=BLUE)
WIDTHS = {"A": 9, "B": 16, "C": 18, "D": 11, "E": 16, "F": 70, "G": 46, "H": 46, "I": 40, "J": 12, "K": 34}


def _f(bold=False, color="000000", size=10):
    return Font(name="Calibri", bold=bold, color=color, size=size)


def build_cases_workbook(scenarios, product: str, source_name: str = "", prepared_by: str = "", notes=None) -> bytes:
    """`scenarios` = the ticked scenarios, already in the order they should appear."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    heads = {"A": "Test Case ID", "B": "Service Type", "C": "Order Type", "D": "IT System", "E": "Product", "F": "Test Scenario",
             "G": "Expected result in CRM/OSS/BSS", "J": "Status", "K": "Workbook reference"}
    for col, text in heads.items():
        c = ws[f"{col}1"]
        c.value = text
    ws.merge_cells("G1:I1")
    for col in "ABCDEFJK":
        ws.merge_cells(f"{col}1:{col}2")
    for col, text in (("G", "CRM"), ("H", "OSS"), ("I", "BSS")):
        ws[f"{col}2"] = text
    for row in (1, 2):
        for col in "ABCDEFGHIJK":
            c = ws[f"{col}{row}"]
            c.font = _f(True, "FFFFFF", 10)
            c.fill = SUB if (row == 2 or col in "GHI") and col in "GHI" else HEAD
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
    for i, sc in enumerate(scenarios, start=1):
        r = i + 2
        vals = [i, product, sc.order_type, sc.it_system, product, sc.text, sc.crm, sc.oss, sc.bss, "Not Run", sc.basis]
        for col, v in zip("ABCDEFGHIJK", vals):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = _f(size=10, bold=(col == "A"), color=("63758B" if col == "K" else "000000"))
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if col in "ACDJ" else "left")
    last = len(scenarios) + 2
    for r in range(3, last + 1):                       # explicit heights: wrapped text does not always auto-fit on open
        lines = 1
        for col in "FGHIK":
            text = str(ws[f"{col}{r}"].value or "")
            per_line = max(8, int(WIDTHS[col] * 1.1))
            lines = max(lines, sum(max(1, -(-len(part) // per_line)) for part in text.split("\n")))
        ws.row_dimensions[r].height = min(400, max(30, lines * 13 + 6))
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20
    if scenarios:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"J3:J{last}")
        ws.conditional_formatting.add(f"J3:J{last}", CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", bgColor="C6F0DF")))
        ws.conditional_formatting.add(f"J3:J{last}", CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", bgColor="FFD0CA")))
        ws.conditional_formatting.add(f"J3:J{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", bgColor="FFE9A8")))
        ws.auto_filter.ref = f"A2:K{last}"
    ws.freeze_panes = "A3"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = BLUE

    # ----------------------------------------------------------------- summary
    sm = wb.create_sheet("Summary")
    sm["A1"] = "TestX · UAT Test Cases"
    sm["A1"].font = _f(True, NAVY, 18)
    sm["A2"] = product
    sm["A2"].font = _f(True, BLUE, 12)
    meta = [("Generated", datetime.datetime.now().strftime("%d.%m.%Y %H:%M")), ("Source workbook", source_name or "-"),
            ("Prepared by", prepared_by or "-"), ("Test cases", len(scenarios))]
    for i, (k, v) in enumerate(meta, start=4):
        sm.cell(i, 1, k).font = _f(True, NAVY)
        sm.cell(i, 2, v).font = _f()
        sm.cell(i, 2).alignment = Alignment(horizontal="left")
    r0 = 4 + len(meta) + 1
    for title, key in (("By order type", lambda s: s.order_type), ("By IT system", lambda s: s.it_system),
                       ("By area", lambda s: s.category)):
        sm.cell(r0, 1, title).font = _f(True, "FFFFFF")
        sm.cell(r0, 2, "Test cases").font = _f(True, "FFFFFF")
        for col in (1, 2):
            sm.cell(r0, col).fill = HEAD
            sm.cell(r0, col).border = BORDER
        counts: dict = {}
        for s in scenarios:
            counts[key(s)] = counts.get(key(s), 0) + 1
        for j, (k, n) in enumerate(counts.items(), start=r0 + 1):
            sm.cell(j, 1, k).font = _f()
            sm.cell(j, 2, n).font = _f()
            sm.cell(j, 2).alignment = Alignment(horizontal="left")
            for col in (1, 2):
                sm.cell(j, col).border = BORDER
        r0 += len(counts) + 2
    if notes:
        sm.cell(r0, 1, "Notes").font = _f(True, NAVY, 11)
        for j, n in enumerate(notes, start=r0 + 1):
            sm.cell(j, 1, n).font = _f()
            sm.cell(j, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sm.merge_cells(start_row=j, start_column=1, end_row=j, end_column=4)
            sm.row_dimensions[j].height = max(15, 15 * (len(n) // 70 + 1))
    sm.column_dimensions["A"].width = 30
    sm.column_dimensions["B"].width = 26
    sm.column_dimensions["C"].width = 20
    sm.column_dimensions["D"].width = 20
    sm.sheet_view.showGridLines = False
    sm.sheet_properties.tabColor = NAVY
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
