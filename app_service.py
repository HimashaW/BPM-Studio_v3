"""Generation service behind the web UI (no Streamlit imports - testable on its own).

Takes the uploaded Excel bytes + the form values and returns everything the
user can download: the Word guide, the flowchart images, the editable
PowerPoint and a combined ZIP.
"""
from __future__ import annotations

import io
import re
import tempfile
import threading
import zipfile
import datetime
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

import pipeline_core as core
from render.task_flow_image import render_flow_images

TEMPLATE_PATH = str(Path(__file__).parent / "assets" / "process_guide_template.docx")
_LOCK = threading.Lock()          # matplotlib / shared temp dirs are not thread-safe

# sheets the pipeline reads (prefix -> friendly name)
EXPECTED_SHEETS = [
    ("1.", "Product Description"),
    ("2.", "Product Hierarchy"),
    ("3.", "Special Requirements"),
    ("5.1", "Order Scenarios"),
    ("6.3", "OSS SO Task List"),
    ("6.4", "Task Descriptions"),
]


@dataclass
class GenParams:
    product: str
    sheet: str
    doc_no: str = ""
    issue_no: str = "0"
    issued_by: str = "Business Process Management Section"
    revision_no: str = "0"
    revision_date: str = ""
    export_pptx: bool = True
    gemini_key: str | None = None
    contacts: dict = field(default_factory=dict)


@dataclass
class GenOutput:
    docx_name: str
    docx_bytes: bytes
    pptx_name: str | None
    pptx_bytes: bytes | None
    flow_zip_name: str
    flow_zip_bytes: bytes
    all_zip_name: str
    all_zip_bytes: bytes
    flow_images: dict            # order type -> [png bytes]
    order_types: list            # [{"name":..., "tasks":..., "images":...}]
    warnings: list
    used_ai: bool


def inspect_workbook(excel_bytes: bytes) -> dict:
    """Sheet names, the detected task-list sheet and which expected sheets exist."""
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    cands = [s for s in names if s.strip().startswith("6.3") and "task" in s.lower()] \
        or [s for s in names if s.strip().startswith("6.3")]
    checks = []
    for prefix, label in EXPECTED_SHEETS:
        hit = next((s for s in names if s.strip().startswith(prefix)), None)
        checks.append({"label": label, "prefix": prefix, "sheet": hit, "found": hit is not None})
    return {"sheets": names, "detected": cands[0] if cands else None, "checks": checks}


def safe_name(text: str, default="Product") -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", text).strip("_") or default


def default_product_name(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"(?i)^product[_ ]configuration[_ ]document[_ ]*v?[_ ]*\d+(?:[._]\d+)*[_ ]*", "", stem)
    stem = re.sub(r"(?i)[_ ]*\d{1,2}[A-Za-z]{3}\d{4}.*$", "", stem)
    return stem.replace("_", " ").strip(" -")


def generate(excel_bytes: bytes, p: GenParams, progress=None) -> GenOutput:
    say = progress or (lambda m: None)
    product = p.product.strip()
    safe = safe_name(product)
    with _LOCK, tempfile.TemporaryDirectory(prefix="sdg_") as td:
        td = Path(td)
        xlsx = td / "input.xlsx"
        xlsx.write_bytes(excel_bytes)
        flow_dir = td / "flowcharts"
        flow_dir.mkdir()

        used_ai = False
        try:
            result = core.extract_all(
                str(xlsx), product=product, sheet=p.sheet, progress_cb=say,
                ai_enhance=bool(p.gemini_key), gemini_api_key=p.gemini_key or None,
            )
            used_ai = bool(p.gemini_key)
        except Exception as exc:
            if not p.gemini_key:
                raise
            say(f"AI introduction unavailable ({type(exc).__name__}) - continuing without it.")
            result = core.extract_all(str(xlsx), product=product, sheet=p.sheet, progress_cb=say)
        core.normalize_attribute_names(result)

        say("Rendering the Word guide ...")
        docx_path = td / f"{safe}_Service_Delivery_Guide.docx"
        today = datetime.date.today().strftime("%d.%m.%Y")
        core.render(
            result, str(docx_path),
            doc_version="1.0", doc_no=p.doc_no or "[TBD]", issue_no=p.issue_no or "0",
            issue_date=today, revision_no=p.revision_no or "0",
            revision_date=p.revision_date or today,
            issued_by=p.issued_by or "Business Process Management Section",
            template_path=TEMPLATE_PATH, include_task_flowchart=True,
            **{k: v for k, v in (p.contacts or {}).items() if v},
        )

        say("Drawing the swimlane flowcharts ...")
        flow_images, order_info = {}, []
        for sec in result.order_types:
            paths, _w = render_flow_images(sec.tasks, str(flow_dir), sec.order_type, product)
            flow_images[sec.order_type] = [Path(x).read_bytes() for x in paths]
            order_info.append({"name": sec.order_type, "tasks": sec.task_count, "images": len(paths)})

        pptx_name = pptx_bytes = None
        if p.export_pptx:
            say("Building the editable PowerPoint ...")
            pptx_name = f"{safe}_Flowcharts_Editable.pptx"
            try:
                core.export_flowcharts_pptx(result, str(td / pptx_name), product)
                pptx_bytes = (td / pptx_name).read_bytes()
            except ValueError:
                pptx_name = None
                result.warnings.append("PowerPoint skipped: no eligible tasks for any flowchart.")

        flow_zip = io.BytesIO()
        with zipfile.ZipFile(flow_zip, "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(flow_dir.glob("*.png")):
                z.write(f, f.name)
        docx_bytes = docx_path.read_bytes()
        all_zip = io.BytesIO()
        with zipfile.ZipFile(all_zip, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(docx_path.name, docx_bytes)
            if pptx_bytes:
                z.writestr(pptx_name, pptx_bytes)
            for f in sorted(flow_dir.glob("*.png")):
                z.write(f, f"flowcharts/{f.name}")

    return GenOutput(
        docx_name=docx_path.name, docx_bytes=docx_bytes, pptx_name=pptx_name, pptx_bytes=pptx_bytes,
        flow_zip_name=f"{safe}_Flowchart_Images.zip", flow_zip_bytes=flow_zip.getvalue(),
        all_zip_name=f"{safe}_SDG_Package.zip", all_zip_bytes=all_zip.getvalue(),
        flow_images=flow_images, order_types=order_info, warnings=list(result.warnings), used_ai=used_ai,
    )


# ---------------------------------------------------------------------------
# Read-only preview used by the "Review extracted information" step of the UI.
# It only READS the workbook (same extractors as generate()); it never writes
# documents and does not change how generate() works.
# ---------------------------------------------------------------------------
def preview_workbook(excel_bytes: bytes, sheet: str, product: str = "") -> dict:
    from render import task_flow as tf
    with tempfile.TemporaryDirectory(prefix="sdg_prev_") as td:
        xlsx = Path(td) / "input.xlsx"
        xlsx.write_bytes(excel_bytes)
        result = core.extract_all(str(xlsx), product=product or "Product", sheet=sheet, progress_cb=None)
    flow_warnings, order_types = [], []
    for sec in result.order_types:
        ft, w = tf.prepare_flow_tasks_ex(sec.tasks)
        flow_warnings += [f"[{sec.order_type}] {x}" for x in w]
        lanes = []
        for t in ft:
            if t.lane not in lanes:
                lanes.append(t.lane)
        order_types.append({
            "name": sec.order_type, "tasks": sec.task_count, "in_flowchart": len(ft),
            "lanes": len(lanes), "decisions": sum(1 for t in ft if t.condition),
        })
    return {
        "order_types": order_types,
        "total_tasks": sum(o["tasks"] for o in order_types),
        "hierarchy": [{"kind": n.kind, "name": n.name, "attributes": len(n.attributes)} for n in result.hierarchy],
        "business_rules": len(result.business_rule_blocks),
        "business_rule_tables": sum(1 for b in result.business_rule_blocks if b.tables),
        "order_scenarios": len(result.order_scenarios),
        "intro_paragraphs": len(result.introduction_text),
        "warnings": list(result.warnings) + flow_warnings,
    }
