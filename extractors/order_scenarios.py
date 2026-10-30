"""Extractor for Section 3 order-scenario data from sheet 5.1 MDM_Order Scenarios.

Expected source columns:
    Service Type | Scenario ID | Order Sub Type | Condition |
    Root Product Action | Order Type | Comments

The document's Section 3 table is mapped as:
    Business Scenario   <- Service Type
    CRM Order Sub Type  <- Order Sub Type
    CRM Order Condition <- Condition
    OSS Service Order   <- Order Type

Rows without an Order Type are ignored because the document table is intended
for order types such as CREATE, DELETE, etc.
"""
from __future__ import annotations

import openpyxl
from extractors.sheet_utils import find_sheet


def _norm(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split()).lower()


def _display(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


ALIASES = {
    "service_type": {"service type", "servicetype"},
    "scenario_id": {"scenario id", "scenarioid"},
    "order_sub_type": {"order sub type", "ordersubtype"},
    "condition": {"condition", "order condition", "ordercondition"},
    "root_product_action": {"root product action", "rootproductaction"},
    "order_type": {"order type", "ordertype"},
    "comments": {"comments", "comment"},
}


def _find_header(ws, max_rows: int = 100):
    for r in range(1, min(ws.max_row, max_rows) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            value = _norm(ws.cell(r, c).value)
            for field, aliases in ALIASES.items():
                if value in aliases:
                    found[field] = c
        required = {"service_type", "order_sub_type", "condition", "order_type"}
        if required.issubset(found):
            return r, found
    return None


def extract_order_scenarios(workbook_path: str, sheet_name: str = "5.1 MDM_Order Scenarios"):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    resolved = find_sheet(wb, sheet_name)
    ws = wb[resolved]

    header = _find_header(ws)
    if header is None:
        raise ValueError(
            f"Could not find the required columns in '{resolved}'. "
            "Expected at least: Service Type, Order Sub Type, Condition, Order Type."
        )

    header_row, cols = header
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        order_type = _display(ws.cell(r, cols["order_type"]).value)
        if not order_type:
            continue

        rows.append({
            "business_scenario": _display(ws.cell(r, cols["service_type"]).value),
            "crm_order_sub_type": _display(ws.cell(r, cols["order_sub_type"]).value),
            "crm_order_condition": _display(ws.cell(r, cols["condition"]).value),
            "oss_service_order": order_type,
        })

    return rows
