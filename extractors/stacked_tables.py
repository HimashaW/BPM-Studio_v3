"""
Extractor for the "stacked tables" pattern: a single sheet containing
multiple tables back to back, each announced by a WIDE MERGED title row
(e.g. "Comprehensive Task List for EdgeDefend - CREATE Order Type"),
followed by a header row, followed by data rows, ended by a blank-row run
or the next title row.

This pattern was identified in sheet "6.3 OSS SO Task List". Reused here
as a general-purpose extractor in case other sheets share the pattern
(e.g. 6.1 / 6.4 if they turn out to be structured similarly).

Anchoring on merged cells rather than fixed row numbers means the
extractor keeps working if rows are inserted/removed upstream.
"""
from __future__ import annotations
import re
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from models.schemas import Task, OrderTypeSection


def _is_blank_row(vals: list) -> bool:
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in vals)


def _find_title_rows(ws: Worksheet, title_keyword: str, min_merge_cols: int = 5) -> dict[int, str]:
    """Locate section title rows via wide merged cells containing `title_keyword`."""
    title_rows: dict[int, str] = {}
    for rng in ws.merged_cells.ranges:
        if (rng.max_col - rng.min_col) >= min_merge_cols:
            val = ws.cell(row=rng.min_row, column=rng.min_col).value
            if val and title_keyword in str(val):
                title_rows[rng.min_row] = str(val).strip()

    # Row 1 title may not be merged (observed in 6.3) - handle separately
    r1_val = ws.cell(row=1, column=1).value
    if r1_val and title_keyword in str(r1_val) and 1 not in title_rows:
        title_rows[1] = str(r1_val).strip()

    return title_rows


# Column layout observed in sheet 6.3. If another sheet uses this pattern
# with different columns, add a new COLS map and pass it in.
DEFAULT_COLS = {
    1: "seq",
    2: "task_name",
    3: "duration",
    4: "cancellation_wo",
    5: "it_system",
    6: "work_group",
    7: "task_condition",
    8: "mandatory_attributes",
    9: "special_configurations",
}


def extract_order_types(
    workbook_path: str,
    sheet_name: str,
    product: str,
    title_keyword: str = "Order Type",
    group_extract_pattern: str = r"-\s*([A-Z\-]+)\s*Order Type",
    cols: dict[int, str] | None = None,
    warnings: list[str] | None = None,
) -> list[OrderTypeSection]:
    """
    Parse a stacked-tables sheet into a list of OrderTypeSection.

    `product` is the product/service name this workbook describes (e.g.
    "CYS_EdgeDefend") - stamped onto every section so the renderer can
    build product-specific headings without hardcoding a name anywhere.

    Raises ValueError if no title rows are found - fail loudly rather than
    silently returning an empty result, since that usually means the
    sheet's structure changed upstream.

    Row-level data-quality issues (e.g. a text value in a numeric column)
    are NOT fatal - they're coerced to None and appended to `warnings`
    (if a list is provided) so they surface in the pipeline report instead
    of crashing the whole run.
    """
    cols = cols or DEFAULT_COLS
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}"
        )
    ws = wb[sheet_name]

    # (row, col) -> value of the top-left cell, for every vertically merged range
    merged_top = {}
    for mr in ws.merged_cells.ranges:
        top_val = ws.cell(mr.min_row, mr.min_col).value
        for rr in range(mr.min_row + 1, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                merged_top[(rr, cc)] = top_val

    title_rows = _find_title_rows(ws, title_keyword)
    if not title_rows:
        raise ValueError(
            f"No section title rows found in '{sheet_name}' matching "
            f"keyword '{title_keyword}'. Sheet structure may have changed."
        )

    sorted_titles = sorted(title_rows.keys())
    sections: list[OrderTypeSection] = []

    for i, title_row in enumerate(sorted_titles):
        title = title_rows[title_row]
        m = re.search(group_extract_pattern, title)
        order_type = m.group(1).strip() if m else title

        header_row = title_row + 1
        data_start = header_row + 1
        next_title_row = sorted_titles[i + 1] if i + 1 < len(sorted_titles) else ws.max_row + 1

        tasks: list[Task] = []
        r = data_start
        blank_streak = 0
        while r < next_title_row:
            vals = [ws.cell(row=r, column=c).value for c in range(1, max(cols) + 1)]
            if _is_blank_row(vals):
                blank_streak += 1
                if blank_streak >= 2:
                    break
                r += 1
                continue
            blank_streak = 0

            field_kwargs = {}
            for col_idx, field_name in cols.items():
                v = vals[col_idx - 1]
                if isinstance(v, str):
                    v = v.strip()
                if field_name == "seq" and v is not None:
                    # Excel stores some seq values as text ("00-1") and others
                    # as numbers (10) - normalize all to string for consistency.
                    v = str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
                elif field_name == "duration" and v is not None:
                    # duration should be numeric; if it isn't (data-entry
                    # error upstream, e.g. "NO" typed into the wrong column),
                    # don't crash the whole extraction - drop it and warn.
                    try:
                        v = float(v)
                    except (TypeError, ValueError):
                        if warnings is not None:
                            warnings.append(
                                f"{sheet_name}!row{r}: non-numeric value "
                                f"'{v}' found in duration column - set to "
                                f"None. Check for a shifted/misaligned row "
                                f"in the source sheet."
                            )
                        v = None
                elif field_name not in ("seq", "duration") and v is not None and not isinstance(v, str):
                    # Every other column is expected to hold text (e.g.
                    # "YES"/"NO", a system name, a free-text condition).
                    # If Excel has a raw number there instead (data-entry
                    # slip, or a differently-shaped row), coerce to string
                    # rather than letting the schema reject the whole row.
                    if warnings is not None:
                        warnings.append(
                            f"{sheet_name}!row{r}: expected text in "
                            f"'{field_name}' column but found a number "
                            f"({v!r}) - converted to text. Check this row "
                            f"isn't shifted/misaligned in the source sheet."
                        )
                    v = str(v)
                field_kwargs[field_name] = v
                if field_name == "seq" and v is None and (r, col_idx) in merged_top:
                    inherited = merged_top[(r, col_idx)]
                    if inherited is not None:
                        field_kwargs["flow_seq"] = (
                            str(int(inherited)) if isinstance(inherited, float) and inherited.is_integer()
                            else str(inherited).strip()
                        )

            # task_name is required by schema - skip rows that don't have one
            if field_kwargs.get("task_name"):
                tasks.append(Task(
                    **field_kwargs,
                    source_sheet=sheet_name,
                    source_row=r,
                ))
            r += 1

        sections.append(OrderTypeSection(
            order_type=order_type,
            product=product,
            title=title,
            source_sheet=sheet_name,
            source_rows=f"{title_row}-{r - 1}",
            tasks=tasks,
        ))

    return sections
