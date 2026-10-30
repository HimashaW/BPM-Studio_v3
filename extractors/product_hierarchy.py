"""
Extractor for sheet "2. Product Hierarchy".

Per spec:
  - Column C ("Entity Type") marks each row's kind: "Product" or "Product
    Component" anchors a new hierarchy node; column D ("Entity Name") on
    that same row is the node's display name (e.g. C6="Product",
    D6="CYS_EdgeDefend").
  - Every row immediately below an anchor row where Entity Type =
    "Attribute" belongs to that anchor, until the next "Product"/"Product
    Component" anchor row is reached. Those rows' "Entity Name",
    "Mandatory?", "Comments" columns become that node's attribute table
    (rendered as 2.1.x.1 Attribute Definitions).
  - Every row where Entity Type = "Value", following an "Attribute" row
    (until the next Attribute/Product/Product Component row), belongs to
    that most-recent attribute. Its "Entity Name" column holds one
    allowed value for that attribute (rendered as 2.1.x.2 Attribute List
    of Values, with the Attribute cell merged across all its values).

ASSUMED STRUCTURE (not yet fully validated against a real sheet): a
header row somewhere near the top containing "Entity Name", "Mandatory?",
"Comments" column labels - located by scanning, not a fixed row number,
since the kind/name columns (C/D) are given as fixed but the attribute
detail columns are not.
"""
from __future__ import annotations
import openpyxl
from models.schemas import Attribute, ProductComponent
from extractors.sheet_utils import find_sheet

KIND_COL = 3   # column C ("Entity Type")
NAME_COL = 4   # column D ("Entity Name")
ANCHOR_VALUES = {"product", "product component"}
ATTRIBUTE_MARKER = "attribute"
VALUE_MARKER = "value"

ATTRIBUTE_DETAIL_COLUMNS = ["Entity Name", "Mandatory?", "Comments"]
MAX_HEADER_SEARCH_ROWS = 15


def _find_attribute_detail_columns(ws) -> dict[str, int]:
    """Locate 'Entity Name' / 'Mandatory?' / 'Comments' column indices by
    scanning the first few rows, rather than assuming fixed positions."""
    found: dict[str, int] = {}
    for r in range(1, min(MAX_HEADER_SEARCH_ROWS, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is None:
                continue
            val_norm = str(val).strip().lower()
            for target in ATTRIBUTE_DETAIL_COLUMNS:
                if val_norm == target.lower() and target not in found:
                    found[target] = c
        if len(found) == len(ATTRIBUTE_DETAIL_COLUMNS):
            break
    return found


def extract_hierarchy(workbook_path: str, sheet_name: str = "2. Product Hierarchy") -> list[ProductComponent]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    resolved = find_sheet(wb, sheet_name)
    ws = wb[resolved]

    detail_cols = _find_attribute_detail_columns(ws)
    missing = set(ATTRIBUTE_DETAIL_COLUMNS) - detail_cols.keys()
    if missing:
        raise ValueError(
            f"Could not find column(s) {missing} in the first "
            f"{MAX_HEADER_SEARCH_ROWS} rows of '{sheet_name}'. Check exact "
            f"header text in the source sheet."
        )

    # 1. find all anchor rows (Product / Product Component)
    anchors: list[tuple[int, str, str]] = []  # (row, kind, name)
    for r in range(1, ws.max_row + 1):
        kind_val = ws.cell(row=r, column=KIND_COL).value
        if kind_val and str(kind_val).strip().lower() in ANCHOR_VALUES:
            name_val = ws.cell(row=r, column=NAME_COL).value
            anchors.append((r, str(kind_val).strip(), str(name_val).strip() if name_val else "(unnamed)"))

    if not anchors:
        raise ValueError(
            f"No rows with 'Product' or 'Product Component' found in column "
            f"C of '{sheet_name}'. Check the sheet structure/column letters."
        )

    # 2. for each anchor, collect Attribute rows (and their following
    #    Value rows) until the next anchor
    components: list[ProductComponent] = []
    for i, (anchor_row, kind, name) in enumerate(anchors):
        next_anchor_row = anchors[i + 1][0] if i + 1 < len(anchors) else ws.max_row + 1
        attributes: list[Attribute] = []
        current_attr: Attribute | None = None
        for r in range(anchor_row + 1, next_anchor_row):
            kind_val = ws.cell(row=r, column=KIND_COL).value
            kind_norm = str(kind_val).strip().lower() if kind_val else None

            if kind_norm == ATTRIBUTE_MARKER:
                entity_name = ws.cell(row=r, column=detail_cols["Entity Name"]).value
                mandatory = ws.cell(row=r, column=detail_cols["Mandatory?"]).value
                comments = ws.cell(row=r, column=detail_cols["Comments"]).value
                if entity_name and str(entity_name).strip():
                    attr_name = str(entity_name).strip()
                    # Some source workbooks repeat the same Attribute row once
                    # per LOV value. Consolidate those rows into one logical
                    # attribute so the Word LOV table shows the attribute once.
                    existing = next(
                        (a for a in attributes if a.entity_name.strip().casefold() == attr_name.casefold()),
                        None,
                    )
                    if existing is not None:
                        if not existing.mandatory and mandatory:
                            existing.mandatory = str(mandatory).strip()
                        if not existing.comments and comments:
                            existing.comments = str(comments).strip()
                        current_attr = existing
                    else:
                        current_attr = Attribute(
                            entity_name=attr_name,
                            mandatory=str(mandatory).strip() if mandatory else None,
                            comments=str(comments).strip() if comments else None,
                            source_sheet=sheet_name,
                            source_row=r,
                        )
                        attributes.append(current_attr)
                else:
                    current_attr = None

            elif kind_norm == VALUE_MARKER:
                if current_attr is not None:
                    val = ws.cell(row=r, column=detail_cols["Entity Name"]).value
                    if val and str(val).strip():
                        display_val = str(val).strip()
                        if display_val not in current_attr.values:
                            current_attr.values.append(display_val)
                # a "Value" row with no preceding attribute in this block
                # is silently skipped - orphaned value, nothing to attach it to

            elif kind_norm in ANCHOR_VALUES:
                # shouldn't normally hit this (loop bound stops at next
                # anchor), but guards against a mis-anchored row
                current_attr = None

        components.append(ProductComponent(
            kind=kind, name=name, attributes=attributes,
            source_sheet=sheet_name, source_row=anchor_row,
        ))

    return components
