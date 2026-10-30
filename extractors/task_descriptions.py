"""Extractor for task descriptions from the 6.4 work-order description sheet.

Expected columns:
    Order Type | Task Name | Work Order Description

The extractor is deliberately header-driven so minor column-position changes do
not break the pipeline.  Descriptions are keyed by (Order Type, Task Name),
with a task-name-only fallback when the order type is blank in the source.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

import openpyxl


def _norm(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split()).lower()


def _display(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_64_sheet(wb):
    # Prefer a sheet whose name starts with 6.4; this is more robust than
    # requiring an exact title because real workbooks often have extra text.
    candidates = [name for name in wb.sheetnames if name.strip().startswith("6.4")]
    if not candidates:
        raise ValueError(
            "Could not find a 6.4 sheet for task descriptions. "
            f"Available sheets: {wb.sheetnames}"
        )
    return candidates[0]


def _find_header_row(ws, max_scan_rows: int = 100) -> Optional[tuple[int, dict[str, int]]]:
    aliases = {
        "order_type": {"order type", "ordertype"},
        "task_name": {"task name", "taskname"},
        "description": {
            "work order description",
            "workorder description",
            "work order description <br>",
            "description",
        },
    }

    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        found: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            value = _norm(ws.cell(r, c).value)
            for field, names in aliases.items():
                if value in names:
                    found[field] = c
        if {"order_type", "task_name", "description"}.issubset(found):
            return r, found
    return None


def extract_task_descriptions(workbook_path: str) -> dict[tuple[str, str], str]:
    """Return {(order_type, task_name): description} from the 6.4 sheet."""
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet_name = _find_64_sheet(wb)
    ws = wb[sheet_name]

    header = _find_header_row(ws)
    if header is None:
        raise ValueError(
            f"Could not find the columns 'Order Type', 'Task Name' and "
            f"'Work Order Description' in sheet '{sheet_name}'."
        )

    header_row, cols = header
    descriptions: dict[tuple[str, str], str] = {}

    for r in range(header_row + 1, ws.max_row + 1):
        order_type = _display(ws.cell(r, cols["order_type"]).value)
        task_name = _display(ws.cell(r, cols["task_name"]).value)
        description = _display(ws.cell(r, cols["description"]).value)
        if not task_name:
            continue
        descriptions[(_norm(order_type), _norm(task_name))] = description

    return descriptions


def attach_task_descriptions(order_types, descriptions: dict[tuple[str, str], str]):
    """Populate Task.description in-place and return the number matched."""
    matched = 0
    by_task_name: dict[str, list[str]] = {}
    for (order_type, task_name), description in descriptions.items():
        if task_name:
            by_task_name.setdefault(task_name, []).append(description)

    for section in order_types:
        section_type = _norm(section.order_type)
        for task in section.tasks:
            key = (section_type, _norm(task.task_name))
            description = descriptions.get(key)

            # If 6.4 has no Order Type value for a row, or if the exact pair
            # was not found, use a task-name fallback only when it is
            # unambiguous. This avoids attaching the wrong description to a
            # repeated task name used by different order types.
            if description is None:
                candidates = by_task_name.get(_norm(task.task_name), [])
                if len(candidates) == 1:
                    description = candidates[0]

            if description is not None:
                task.description = description
                matched += 1

    return matched
